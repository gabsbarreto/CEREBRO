from __future__ import annotations

import json
import re
import shutil
import time
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from app.models import JobSettings
from app.services import jobs, projects

STRUCTURED_EXTRACTION_TYPE = "pdf_structured"
STRUCTURED_SHEETS_DIRNAME = "structured_sheets"
SHEET_FILENAME = "sheet.json"
ROWS_JSONL_FILENAME = "rows.jsonl"
ERRORS_JSONL_FILENAME = "errors.jsonl"
RAW_OUTPUT_FILENAME = "structured_raw_response.tsv"
PARSED_ROWS_FILENAME = "structured_parsed_rows.json"
PARSE_ERRORS_FILENAME = "structured_parse_errors.json"
STRUCTURED_EXPORT_SUFFIX = "structured_pdf_export.xlsx"
MAX_ROWS_LIMIT = 500

_BLOCK_SEPARATOR_RE = re.compile(r"(?m)^\s*---\s*$")
_SECTION_RE = re.compile(r"(?m)^(Column name|Question|Rules)\s*###\s*(.*)$")
_SHEET_IMPORT_SECTION_RE = re.compile(
    r"(?mi)^(Sheet name|Context|More information / other preferences|More information|Other preferences|Columns)\s*###\s*(.*)$"
)


class StructuredParseError(ValueError):
    """Raised when model output cannot be parsed as the requested TSV schema."""


def list_sheets(project_id: str) -> list[dict[str, Any]]:
    require_structured_project(project_id)
    root = sheets_root(project_id)
    sheets: list[dict[str, Any]] = []
    for path in sorted(root.glob(f"*/{SHEET_FILENAME}"), key=lambda item: item.parent.name.lower()):
        try:
            sheets.append(read_sheet_file(path))
        except Exception:
            continue
    ensure_sheet_orders(project_id, sheets)
    ensure_unique_sheet_names(project_id, sheets)
    sheets.sort(key=sheet_sort_key)
    return sheets


def create_sheet(
    *,
    project_id: str,
    name: str,
    context: str = "",
    row_unit: str = "",
    columns: list[dict[str, str]] | None = None,
) -> dict[str, Any]:
    require_structured_project(project_id)
    payload = normalized_sheet_payload(
        project_id=project_id,
        sheet_id=unique_sheet_id(project_id, name),
        name=name,
        context=context,
        row_unit=row_unit,
        columns=columns or [],
        created_at=utc_now(),
        sheet_order=next_sheet_order(project_id),
    )
    root = sheet_root(project_id, str(payload["sheet_id"]))
    root.mkdir(parents=True, exist_ok=True)
    jobs.write_json(root / SHEET_FILENAME, payload)
    ensure_sheet_data_files(root)
    return read_sheet_file(root / SHEET_FILENAME)


def update_sheet(
    *,
    project_id: str,
    sheet_id: str,
    name: str,
    context: str = "",
    row_unit: str = "",
    columns: list[dict[str, str]] | None = None,
) -> dict[str, Any]:
    existing = get_sheet(project_id, sheet_id)
    ensure_sheet_editable(existing)
    payload = normalized_sheet_payload(
        project_id=project_id,
        sheet_id=sheet_id,
        name=name,
        context=context,
        row_unit=row_unit,
        columns=columns or [],
        created_at=str(existing.get("created_at") or utc_now()),
        sheet_order=sheet_order(existing),
    )
    jobs.write_json(sheet_root(project_id, sheet_id) / SHEET_FILENAME, payload)
    return read_sheet_file(sheet_root(project_id, sheet_id) / SHEET_FILENAME)


def duplicate_sheet(project_id: str, sheet_id: str) -> dict[str, Any]:
    existing = get_sheet(project_id, sheet_id)
    duplicate_order = sheet_order(existing) + 1
    shift_sheet_orders(project_id, start_order=duplicate_order)
    duplicate_name = unique_duplicate_sheet_name(project_id, str(existing.get("name") or "Sheet"))
    payload = normalized_sheet_payload(
        project_id=project_id,
        sheet_id=unique_sheet_id(project_id, duplicate_name),
        name=duplicate_name,
        context=str(existing.get("context") or ""),
        row_unit=str(existing.get("row_unit") or ""),
        columns=[
            {
                "column_name": str(column.get("column_name") or ""),
                "question": str(column.get("question") or ""),
                "rules": str(column.get("rules") or ""),
            }
            for column in existing.get("columns") or []
        ],
        created_at=utc_now(),
        sheet_order=duplicate_order,
    )
    root = sheet_root(project_id, str(payload["sheet_id"]))
    root.mkdir(parents=True, exist_ok=True)
    jobs.write_json(root / SHEET_FILENAME, payload)
    ensure_sheet_data_files(root)
    return read_sheet_file(root / SHEET_FILENAME)


def duplicate_workbook(project_id: str, sheet_ids: list[str]) -> dict[str, Any]:
    source_project = require_structured_project(project_id)
    selected_ids = {validate_sheet_id(sheet_id) for sheet_id in sheet_ids if str(sheet_id or "").strip()}
    if not selected_ids:
        raise ValueError("Select at least one sheet to duplicate.")

    source_sheets = list_sheets(project_id)
    selected_sheets = [sheet for sheet in source_sheets if str(sheet.get("sheet_id") or "") in selected_ids]
    if len(selected_sheets) != len(selected_ids):
        raise ValueError("One or more selected sheets do not belong to this workbook.")

    project = projects.create_project(
        next_workbook_copy_name(str(source_project.get("name") or "Structured workbook")),
        str(source_project.get("description") or ""),
        STRUCTURED_EXTRACTION_TYPE,
    )
    copied_sheets: list[dict[str, Any]] = []
    try:
        for sheet in selected_sheets:
            copied_sheets.append(
                create_sheet(
                    project_id=str(project["project_id"]),
                    name=str(sheet.get("name") or "Sheet"),
                    context=str(sheet.get("context") or ""),
                    row_unit=str(sheet.get("row_unit") or ""),
                    columns=[
                        {
                            "column_name": str(column.get("column_name") or ""),
                            "question": str(column.get("question") or ""),
                            "rules": str(column.get("rules") or ""),
                        }
                        for column in sheet.get("columns") or []
                    ],
                )
            )
    except Exception:
        shutil.rmtree(projects.project_root(str(project["project_id"])), ignore_errors=True)
        raise
    return {
        "source_project_id": project_id,
        "project": project,
        "sheets": copied_sheets,
    }


def delete_sheet(project_id: str, sheet_id: str) -> None:
    get_sheet(project_id, sheet_id)
    root = sheet_root(project_id, sheet_id)
    if root.exists():
        shutil.rmtree(root)


def get_sheet(project_id: str, sheet_id: str) -> dict[str, Any]:
    require_structured_project(project_id)
    safe_sheet_id = validate_sheet_id(sheet_id)
    path = sheet_root(project_id, safe_sheet_id) / SHEET_FILENAME
    if not path.exists():
        raise FileNotFoundError(f"Structured sheet not found: {safe_sheet_id}")
    return read_sheet_file(path)


def parse_column_blocks(raw_text: str) -> list[dict[str, str]]:
    blocks = [block.strip() for block in _BLOCK_SEPARATOR_RE.split(str(raw_text or "")) if block.strip()]
    if not blocks:
        raise ValueError("Paste at least one column block.")
    columns: list[dict[str, str]] = []
    seen: set[str] = set()
    for index, block in enumerate(blocks, start=1):
        sections = extract_block_sections(block)
        column_name = sections.get("Column name", "").strip()
        question = sections.get("Question", "").strip()
        rules = sections.get("Rules", "").strip()
        if not column_name:
            raise ValueError(f"Column block {index} is missing `Column name ###`.")
        if not question:
            raise ValueError(f"Column block {index} is missing `Question ###`.")
        if column_name in seen:
            raise ValueError(f"Duplicate column name: {column_name}")
        seen.add(column_name)
        columns.append({"column_name": column_name, "question": question, "rules": rules})
    return columns


def parse_sheet_import(raw_text: str, current_sheet: dict[str, Any] | None = None) -> dict[str, Any]:
    text = str(raw_text or "").strip()
    if not text:
        raise ValueError("Paste sheet or column text to import.")
    parsed_fields, column_text = parse_sheet_import_fields(text)
    columns: list[dict[str, str]] = []
    if re.search(r"(?mi)^Column name\s*###", column_text):
        columns = parse_column_blocks(column_text)
    elif re.search(r"(?mi)^Column name\s*###", text):
        columns = parse_column_blocks(text)
    if not parsed_fields and not columns:
        raise ValueError("No sheet fields or column blocks were found.")
    conflicts = import_conflicts(parsed_fields, columns, current_sheet or {})
    return {
        "fields": parsed_fields,
        "columns": columns,
        "count": len(columns),
        "has_conflicts": bool(conflicts["fields"] or conflicts["columns"]),
        "conflicts": conflicts,
    }


def format_sheet_import_text(sheet: dict[str, Any]) -> str:
    columns = normalize_columns(sheet.get("columns") or [], require_questions=False)
    column_blocks = []
    for column in columns:
        column_blocks.append(
            "\n".join(
                [
                    f"Column name ### {column['column_name']}",
                    "",
                    f"Question ### {column.get('question') or ''}",
                    "",
                    "Rules ###",
                    str(column.get("rules") or "").strip(),
                ]
            ).rstrip()
        )
    return (
        "\n".join(
            [
                f"Sheet name ### {str(sheet.get('name') or 'Sheet').strip() or 'Sheet'}",
                "",
                "Context ###",
                str(sheet.get("context") or "").strip(),
                "",
                "More information / other preferences ###",
                str(sheet.get("row_unit") or "").strip(),
                "",
                "Columns ###",
                "",
                "\n\n---\n\n".join(column_blocks),
            ]
        ).rstrip()
        + "\n"
    )


def parse_sheet_import_fields(text: str) -> tuple[dict[str, str], str]:
    matches = list(_SHEET_IMPORT_SECTION_RE.finditer(text))
    if not matches:
        return {}, text
    fields: dict[str, str] = {}
    column_text = ""
    for index, match in enumerate(matches):
        heading = normalize_import_heading(match.group(1))
        inline_value = match.group(2) or ""
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        body = text[start:end]
        value = "\n".join(part for part in [inline_value.strip(), body.strip()] if part).strip()
        if heading == "columns":
            column_text = value
        elif value:
            fields[heading] = value
    if not column_text:
        first_column = re.search(r"(?mi)^Column name\s*###", text)
        if first_column:
            column_text = text[first_column.start() :].strip()
    return fields, column_text


def normalize_import_heading(heading: str) -> str:
    value = re.sub(r"\s+", " ", str(heading or "").strip().lower())
    if value == "sheet name":
        return "name"
    if value == "context":
        return "context"
    if value in {"more information / other preferences", "more information", "other preferences"}:
        return "row_unit"
    if value == "columns":
        return "columns"
    return value


def import_conflicts(
    parsed_fields: dict[str, str],
    parsed_columns: list[dict[str, str]],
    current_sheet: dict[str, Any],
) -> dict[str, list[str]]:
    conflicts = {"fields": [], "columns": []}
    field_labels = {
        "context": "Context",
        "row_unit": "More information / other preferences",
    }
    for key, label in field_labels.items():
        new_value = str(parsed_fields.get(key) or "").strip()
        current_value = str(current_sheet.get(key) or "").strip()
        if new_value and current_value and new_value != current_value:
            conflicts["fields"].append(label)
    existing_names = {
        str(column.get("column_name") or "").strip()
        for column in current_sheet.get("columns") or []
        if str(column.get("column_name") or "").strip()
    }
    for column in parsed_columns:
        column_name = str(column.get("column_name") or "").strip()
        if column_name and column_name in existing_names:
            conflicts["columns"].append(column_name)
    return conflicts


def compile_sheet_prompt(sheet: dict[str, Any]) -> str:
    columns = normalize_columns(sheet.get("columns") or [], require_questions=True)
    header = "\t".join(column["column_name"] for column in columns)
    preferences = (
        str(sheet.get("row_unit") or "").strip()
        or "Use the context and column-specific instructions to decide which rows belong in the output."
    )
    column_sections: list[str] = []
    for column in columns:
        rules = str(column.get("rules") or "").strip() or "No additional rules."
        column_sections.append(
            "\n".join(
                [
                    f"Column name ### {column['column_name']}",
                    "",
                    f"Question ### {column['question']}",
                    "",
                    "Rules ###",
                    "",
                    rules,
                ]
            )
        )
    return "\n\n".join(
        [
            "You are extracting structured data from a scientific PDF.",
            "Context:",
            str(sheet.get("context") or "").strip() or "No additional context provided.",
            "Other preferences:",
            preferences,
            (
                "For each row, answer the following column-specific questions. "
                "The title of each question is the exact output column name. "
                "You must use these column names exactly."
            ),
            "\n\n".join(column_sections),
            "Required output format:",
            "Return the results as tab-separated values.",
            "The first line must contain the column names below, in exactly this order:",
            header,
            "Each following line must contain one corresponding row according to the context, other preferences, and column instructions.",
            "Do not include a Markdown table, code fences, bullets, introductory text, explanations, or any additional lines.",
            "Do not include tabs or line breaks inside any cell.",
            (
                "If multiple values belong in the same cell, keep them in the same cell separated by comma and one space, "
                "unless the column rules say otherwise."
            ),
            "If no eligible data can be extracted, return only the header line.",
        ]
    ).strip()


def parse_tsv_output(raw_output: str, expected_columns: list[str]) -> list[dict[str, str]]:
    columns = [str(column or "").strip() for column in expected_columns if str(column or "").strip()]
    if not columns:
        raise StructuredParseError("Structured sheet has no expected columns.")
    text = str(raw_output or "").strip(" \r\n")
    if not text:
        raise StructuredParseError("Model returned an empty response.")
    text = strip_surrounding_code_fence(text)
    if "```" in text:
        raise StructuredParseError("Model response contains a Markdown code fence.")
    lines = [line.strip(" ") for line in text.splitlines() if line.strip(" ")]
    if not lines:
        raise StructuredParseError("Model returned no non-empty TSV lines.")
    if any(line.startswith("|") for line in lines):
        raise StructuredParseError("Model response appears to be a Markdown table, not TSV.")
    expected_header = "\t".join(columns)
    actual_header = lines[0]
    if actual_header != expected_header:
        raise StructuredParseError(f"TSV header mismatch. Expected `{expected_header}` but got `{actual_header}`.")
    rows: list[dict[str, str]] = []
    for line_number, line in enumerate(lines[1:], start=2):
        cells = line.split("\t")
        if len(cells) != len(columns):
            raise StructuredParseError(
                f"Line {line_number} has {len(cells)} cells but expected {len(columns)}."
            )
        rows.append({column: cells[index].strip() for index, column in enumerate(columns)})
    return rows


def complete_structured_job(
    *,
    job_id: str,
    settings: JobSettings,
    project_id: str | None,
    started_at: float,
    prompt_filename: str,
    prompt_source_path: str,
    system_prompt_file: Path,
    user_prompt_file: Path,
    output_file: Path,
    pdf_path: Path,
) -> None:
    if not project_id:
        raise RuntimeError("Structured PDF jobs require a project id.")
    completed_at = time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())
    duration_seconds = round(time.time() - started_at, 2)
    root = jobs.job_dir(job_id, project_id)
    metadata = jobs.read_metadata(root)
    sheet_id = str(metadata.get("structured_sheet_id") or "")
    sheet = get_sheet(project_id, sheet_id)
    columns = [
        column["column_name"]
        for column in normalize_columns(
            metadata.get("structured_columns") or sheet.get("columns") or [],
            require_questions=False,
        )
    ]
    raw_output = output_file.read_text(encoding="utf-8") if output_file.exists() else ""
    raw_output_path = root / "outputs" / RAW_OUTPUT_FILENAME
    parsed_rows_path = root / "outputs" / PARSED_ROWS_FILENAME
    parse_errors_path = root / "outputs" / PARSE_ERRORS_FILENAME
    raw_output_path.write_text(raw_output, encoding="utf-8")

    row_records: list[dict[str, Any]] = []
    error_records: list[dict[str, Any]] = []
    parse_status = "parsed"
    parse_error = ""
    try:
        parsed_rows = parse_tsv_output(raw_output, columns)
        for index, cells in enumerate(parsed_rows, start=1):
            row_records.append(
                structured_row_record(
                    project_id=project_id,
                    sheet_id=sheet_id,
                    job_id=job_id,
                    metadata=metadata,
                    extracted_at=completed_at,
                    model=settings.rq_screening_model,
                    cells=cells,
                    row_index=index,
                )
            )
    except StructuredParseError as exc:
        parse_status = "failed"
        parse_error = str(exc)
        error_records.append(
            structured_error_record(
                project_id=project_id,
                sheet_id=sheet_id,
                job_id=job_id,
                metadata=metadata,
                extracted_at=completed_at,
                model=settings.rq_screening_model,
                parse_error=parse_error,
                raw_output_file=raw_output_path,
                raw_output=raw_output,
            )
        )

    jobs.write_json(parsed_rows_path, row_records)
    jobs.write_json(parse_errors_path, error_records)
    replace_job_sheet_records(project_id, sheet_id, job_id, row_records, error_records)
    jobs.update_metadata(
        root,
        completed_at=completed_at,
        duration_seconds=duration_seconds,
        **jobs.settings_metadata_updates(
            settings,
            rq_prompt_filename=prompt_filename,
            include_settings=False,
        ),
        rq_prompt_source_path=prompt_source_path,
        rq_system_prompt_file=str(system_prompt_file),
        rq_user_prompt_file=str(user_prompt_file),
        summary_xlsx_path=None,
        pending_rerun_screening_only=False,
        rerun_reuses_ocr=False,
        ocr_complete=settings.openai_input_mode != "pdf_file",
        openai_file_complete=settings.rq_provider == "openai" and settings.openai_input_mode == "pdf_file",
        structured_raw_output_file=str(raw_output_path),
        structured_parsed_rows_file=str(parsed_rows_path),
        structured_parse_errors_file=str(parse_errors_path),
        structured_parse_status=parse_status,
        structured_parse_error=parse_error,
        structured_parsed_row_count=len(row_records),
    )
    message = (
        f"Structured extraction parsed {len(row_records)} row{'' if len(row_records) == 1 else 's'}"
        if parse_status == "parsed"
        else f"Structured extraction completed with parse error: {parse_error}"
    )
    jobs.update_status(
        job_id,
        status="complete",
        stage="complete",
        message=message,
        progress=1.0,
        error=None,
        event={
            "event": "structured_complete",
            "parse_status": parse_status,
            "parsed_rows": len(row_records),
            "parse_error": parse_error,
        },
        project_id=project_id,
    )


def list_structured_jobs(project_id: str) -> list[dict[str, Any]]:
    require_structured_project(project_id)
    records = [
        record
        for record in jobs.list_jobs(limit=0, project_id=project_id)
        if (record.get("metadata") or {}).get("extraction_type") == STRUCTURED_EXTRACTION_TYPE
    ]
    return records


def list_structured_jobs_window(
    project_id: str,
    *,
    sheet_id: str = "",
    offset: int = 0,
    limit: int = 80,
    status: str = "all",
    search: str = "",
) -> dict[str, Any]:
    records = list_structured_jobs(project_id)
    selected_sheet_id = str(sheet_id or "").strip()
    if selected_sheet_id:
        validate_sheet_id(selected_sheet_id)
        records = [
            record
            for record in records
            if str((record.get("metadata") or {}).get("structured_sheet_id") or "") == selected_sheet_id
        ]

    counts = {"all": len(records), "running": 0, "queued": 0, "completed": 0, "failed": 0, "parse_error": 0}
    for record in records:
        key = structured_job_status_key(record)
        counts[key] += 1
        metadata = record.get("metadata") if isinstance(record.get("metadata"), dict) else {}
        if str(metadata.get("structured_parse_status") or "").lower() == "failed":
            counts["parse_error"] += 1

    requested_status = str(status or "all").strip().lower()
    query = str(search or "").strip().lower()
    filtered = [
        record
        for record in records
        if (requested_status == "all" or structured_job_status_key(record) == requested_status)
        and (not query or query in structured_job_search_haystack(record))
    ]
    safe_limit = max(1, min(int(limit or 80), 250))
    safe_offset = max(0, int(offset or 0))
    if filtered and safe_offset >= len(filtered):
        safe_offset = max(0, ((len(filtered) - 1) // safe_limit) * safe_limit)
    return {
        "items": filtered[safe_offset : safe_offset + safe_limit],
        "total": len(filtered),
        "offset": safe_offset,
        "limit": safe_limit,
        "counts": counts,
    }


def structured_job_status_key(record: dict[str, Any]) -> str:
    payload = record.get("status") if isinstance(record.get("status"), dict) else {}
    value = str(payload.get("status") or "queued").strip().lower()
    if value in {"complete", "completed"}:
        return "completed"
    if value in {"failed", "error"}:
        return "failed"
    if value == "running":
        return "running"
    return "queued"


def structured_job_search_haystack(record: dict[str, Any]) -> str:
    metadata = record.get("metadata") if isinstance(record.get("metadata"), dict) else {}
    status = record.get("status") if isinstance(record.get("status"), dict) else {}
    values = [
        record.get("job_id"),
        record.get("filename"),
        status.get("status"),
        status.get("message"),
        metadata.get("original_filename"),
        metadata.get("structured_sheet_name"),
        metadata.get("structured_parse_status"),
        metadata.get("structured_parse_error"),
        metadata.get("rq_screening_model"),
    ]
    return " ".join(str(value or "") for value in values).lower()


def lock_sheet_for_first_run(project_id: str, sheet_id: str, job_id: str) -> dict[str, Any]:
    sheet = get_sheet(project_id, sheet_id)
    if sheet_is_locked(sheet):
        return sheet
    payload = jobs.read_json(sheet_root(project_id, sheet_id) / SHEET_FILENAME)
    now = utc_now()
    payload["locked_at"] = now
    payload["locked_by_job_id"] = str(job_id)
    payload["locked_reason"] = "First extraction run was queued."
    payload["updated_at"] = now
    jobs.write_json(sheet_root(project_id, sheet_id) / SHEET_FILENAME, payload)
    return read_sheet_file(sheet_root(project_id, sheet_id) / SHEET_FILENAME)


def structured_job_result(project_id: str, job_id: str) -> dict[str, Any]:
    require_structured_project(project_id)
    root = jobs.job_dir(job_id, project_id)
    if not root.exists():
        raise FileNotFoundError("Structured PDF job not found.")
    metadata = jobs.read_metadata(root)
    if metadata.get("extraction_type") != STRUCTURED_EXTRACTION_TYPE:
        raise ValueError("Job is not a structured PDF extraction job.")
    return {
        "status": jobs.read_status(job_id, project_id=project_id),
        "metadata": metadata,
        "raw_output": read_text_if_exists(root / "outputs" / RAW_OUTPUT_FILENAME)
        or read_text_if_exists(root / "outputs" / "rq_screening_output.md"),
        "parsed_rows": read_json_if_exists(root / "outputs" / PARSED_ROWS_FILENAME, default=[]),
        "parse_errors": read_json_if_exists(root / "outputs" / PARSE_ERRORS_FILENAME, default=[]),
        "system_prompt": read_text_if_exists(root / "outputs" / "rq_system_prompt.txt"),
        "user_prompt": read_text_if_exists(root / "outputs" / "rq_user_prompt.txt"),
        "job_dir": str(root),
    }


def list_rows(
    *,
    project_id: str,
    sheet_id: str,
    offset: int = 0,
    limit: int = 100,
    search: str = "",
) -> dict[str, Any]:
    sheet = get_sheet(project_id, sheet_id)
    query = str(search or "").strip().lower()
    sheet_records = read_sheet_rows(project_id, sheet_id) + read_sheet_errors(project_id, sheet_id)
    structured_jobs = [
        record
        for record in list_structured_jobs(project_id)
        if str((record.get("metadata") or {}).get("structured_sheet_id") or "") == str(sheet_id)
    ]
    jobs_by_id = {str(record.get("job_id") or ""): record for record in structured_jobs}
    records = [enrich_sheet_record(record, jobs_by_id.get(str(record.get("job_id") or ""))) for record in sheet_records]
    recorded_job_ids = {str(record.get("job_id") or "") for record in records if str(record.get("job_id") or "")}
    for job_record in structured_jobs:
        job_id = str(job_record.get("job_id") or "")
        if job_id in recorded_job_ids:
            continue
        records.append(structured_job_row_record(project_id, sheet_id, job_record))
    records.sort(
        key=lambda item: (
            str(item.get("created_at") or item.get("extracted_at") or item.get("parse_date") or ""),
            str(item.get("source_pdf") or ""),
            int(item.get("row_index") or 0),
            str(item.get("row_id") or item.get("error_id") or item.get("job_row_id") or ""),
        )
    )
    if query:
        records = [record for record in records if query in row_search_haystack(record)]
    safe_offset = max(0, int(offset or 0))
    safe_limit = max(1, min(MAX_ROWS_LIMIT, int(limit or 100)))
    return {
        "sheet": sheet,
        "items": records[safe_offset : safe_offset + safe_limit],
        "total": len(records),
        "offset": safe_offset,
        "limit": safe_limit,
    }


def enrich_sheet_record(record: dict[str, Any], job_record: dict[str, Any] | None) -> dict[str, Any]:
    payload = dict(record)
    metadata = (job_record or {}).get("metadata") or {}
    status = (job_record or {}).get("status") or {}
    payload.setdefault("created_at", str(metadata.get("created_at") or ""))
    payload.setdefault("job_status", str(status.get("status") or "complete"))
    payload.setdefault("status_message", str(status.get("message") or ""))
    payload.setdefault("parse_date", str(payload.get("parse_date") or payload.get("extracted_at") or metadata.get("completed_at") or ""))
    return payload


def structured_job_row_record(project_id: str, sheet_id: str, job_record: dict[str, Any]) -> dict[str, Any]:
    metadata = job_record.get("metadata") or {}
    status = job_record.get("status") or {}
    job_status = str(status.get("status") or "queued")
    parse_status = str(metadata.get("structured_parse_status") or job_status or "queued")
    parse_error = str(metadata.get("structured_parse_error") or status.get("error") or "")
    parse_date = str(metadata.get("completed_at") or "")
    return {
        "job_row_id": f"job-{job_record.get('job_id')}",
        "sheet_id": sheet_id,
        "project_id": project_id,
        "job_id": str(job_record.get("job_id") or ""),
        "source_pdf": str(metadata.get("original_filename") or job_record.get("filename") or ""),
        "source_relative_path": str(metadata.get("source_relative_path") or ""),
        "created_at": str(metadata.get("created_at") or ""),
        "extracted_at": parse_date,
        "parse_date": parse_date,
        "model": str(metadata.get("rq_screening_model") or ""),
        "parse_status": parse_status,
        "parse_error": parse_error,
        "job_status": job_status,
        "status_message": str(status.get("message") or ""),
        "row_index": 0,
        "cells": {},
        "is_job_placeholder": True,
    }


def export_structured_workbook(project_id: str) -> Path:
    project = require_structured_project(project_id)
    sheets = list_sheets(project_id)
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    used_names: set[str] = set()
    if not sheets:
        worksheet = workbook.create_sheet("Structured export")
        worksheet.append(["No structured sheets found"])
    for sheet in sheets:
        worksheet = workbook.create_sheet(safe_worksheet_name(str(sheet.get("name") or "Sheet"), used_names))
        columns = [column["column_name"] for column in normalize_columns(sheet.get("columns") or [], require_questions=False)]
        headers = [
            excel_safe_text(header)
            for header in ["source_pdf", "job_id", "extracted_at", "model", "parse_status", "parse_date", "parse_error"] + columns
        ]
        worksheet.append(headers)
        for record in read_sheet_rows(project_id, str(sheet["sheet_id"])) + read_sheet_errors(project_id, str(sheet["sheet_id"])):
            cells = record.get("cells") if isinstance(record.get("cells"), dict) else {}
            worksheet.append(
                [
                    excel_safe_text(record.get("source_pdf")),
                    excel_safe_text(record.get("job_id")),
                    excel_safe_text(record.get("extracted_at")),
                    excel_safe_text(record.get("model")),
                    excel_safe_text(record.get("parse_status")),
                    excel_safe_text(record.get("parse_date") or record.get("extracted_at")),
                    excel_safe_text(record.get("parse_error")),
                ]
                + [excel_safe_text(cells.get(column)) for column in columns]
            )
        for column in worksheet.columns:
            letter = column[0].column_letter
            worksheet.column_dimensions[letter].width = min(max(len(str(column[0].value or "")) + 4, 14), 60)
    path = projects.project_root(project_id) / f"cerebro_{projects.sanitize_slug(str(project.get('name') or project_id))}_{STRUCTURED_EXPORT_SUFFIX}"
    tmp_path = path.with_name(f".{path.name}.tmp")
    try:
        workbook.save(tmp_path)
        tmp_path.replace(path)
    finally:
        if tmp_path.exists():
            tmp_path.unlink()
    return path


def require_structured_project(project_id: str) -> dict[str, Any]:
    project = projects.get_project(project_id)
    if projects.normalize_extraction_type(project.get("extraction_type")) != STRUCTURED_EXTRACTION_TYPE:
        raise ValueError("Project is not a structured PDF extraction project.")
    return project


def sheets_root(project_id: str) -> Path:
    projects.validate_project_id(project_id)
    return projects.project_root(project_id) / STRUCTURED_SHEETS_DIRNAME


def sheet_root(project_id: str, sheet_id: str) -> Path:
    return sheets_root(project_id) / validate_sheet_id(sheet_id)


def validate_sheet_id(sheet_id: str) -> str:
    safe_id = str(sheet_id or "").strip()
    if not safe_id:
        raise ValueError("Sheet id is required.")
    projects.validate_project_id(safe_id)
    return safe_id


def unique_sheet_id(project_id: str, name: str) -> str:
    stem = projects.sanitize_slug(name)
    for _attempt in range(20):
        candidate = f"{stem}-{uuid.uuid4().hex[:8]}"
        if not sheet_root(project_id, candidate).exists():
            return candidate
    return uuid.uuid4().hex


def unique_duplicate_sheet_name(project_id: str, name: str) -> str:
    existing = {str(sheet.get("name") or "") for sheet in list_sheets(project_id)}
    base = duplicate_base_name(name)
    return next_duplicate_name(base, existing)


def next_workbook_copy_name(name: str) -> str:
    base = duplicate_base_name(name or "Structured workbook")
    existing = {str(project.get("name") or "").casefold() for project in projects.list_projects()}
    index = 1
    while f"{base} ({index})".casefold() in existing:
        index += 1
    return f"{base} ({index})"


def duplicate_base_name(name: str) -> str:
    return re.sub(r"\s+\(\d+\)$", "", str(name or "Sheet").strip() or "Sheet").strip() or "Sheet"


def next_duplicate_name(base: str, existing: set[str]) -> str:
    clean_base = str(base or "Sheet").strip() or "Sheet"
    index = 1
    while f"{clean_base} ({index})" in existing:
        index += 1
    return f"{clean_base} ({index})"


def sheet_sort_key(sheet: dict[str, Any]) -> tuple[int, str, str]:
    return (sheet_order(sheet), str(sheet.get("created_at") or ""), str(sheet.get("sheet_id") or ""))


def sheet_order(sheet: dict[str, Any]) -> int:
    try:
        return int(sheet.get("sheet_order"))
    except (TypeError, ValueError):
        return 1_000_000


def ensure_sheet_orders(project_id: str, sheets: list[dict[str, Any]]) -> None:
    if not any(sheet_order(sheet) >= 1_000_000 for sheet in sheets):
        return
    ordered = sorted(sheets, key=lambda sheet: (str(sheet.get("created_at") or ""), str(sheet.get("sheet_id") or "")))
    for index, sheet in enumerate(ordered):
        if sheet_order(sheet) == index:
            continue
        sheet["sheet_order"] = index
        jobs.write_json(sheet_root(project_id, str(sheet["sheet_id"])) / SHEET_FILENAME, serialized_sheet_payload(sheet))


def ensure_unique_sheet_names(project_id: str, sheets: list[dict[str, Any]]) -> None:
    used_names: set[str] = set()
    for sheet in sorted(sheets, key=sheet_sort_key):
        name = str(sheet.get("name") or "Sheet").strip() or "Sheet"
        if name not in used_names:
            if sheet.get("name") != name:
                sheet["name"] = name
                jobs.write_json(sheet_root(project_id, str(sheet["sheet_id"])) / SHEET_FILENAME, serialized_sheet_payload(sheet))
            used_names.add(name)
            continue
        base = duplicate_base_name(name)
        candidate = next_duplicate_name(base, used_names)
        sheet["name"] = candidate
        used_names.add(candidate)
        jobs.write_json(sheet_root(project_id, str(sheet["sheet_id"])) / SHEET_FILENAME, serialized_sheet_payload(sheet))


def next_sheet_order(project_id: str) -> int:
    orders = [sheet_order(sheet) for sheet in list_sheets(project_id)]
    real_orders = [order for order in orders if order < 1_000_000]
    if real_orders:
        return max(real_orders) + 1
    return len(orders)


def shift_sheet_orders(project_id: str, *, start_order: int) -> None:
    for sheet in reversed(list_sheets(project_id)):
        order = sheet_order(sheet)
        if order >= start_order:
            sheet["sheet_order"] = order + 1
            jobs.write_json(sheet_root(project_id, str(sheet["sheet_id"])) / SHEET_FILENAME, serialized_sheet_payload(sheet))


def serialized_sheet_payload(sheet: dict[str, Any]) -> dict[str, Any]:
    payload = {
        "sheet_id": validate_sheet_id(str(sheet.get("sheet_id") or "")),
        "project_id": projects.validate_project_id(str(sheet.get("project_id") or "")),
        "name": str(sheet.get("name") or "Sheet"),
        "context": str(sheet.get("context") or ""),
        "row_unit": str(sheet.get("row_unit") or ""),
        "columns": normalize_columns(sheet.get("columns") or [], require_questions=False),
        "sheet_order": sheet_order(sheet),
        "created_at": str(sheet.get("created_at") or ""),
        "updated_at": str(sheet.get("updated_at") or ""),
        "locked_at": str(sheet.get("locked_at") or ""),
        "locked_by_job_id": str(sheet.get("locked_by_job_id") or ""),
        "locked_reason": str(sheet.get("locked_reason") or ""),
    }
    return payload


def normalized_sheet_payload(
    *,
    project_id: str,
    sheet_id: str,
    name: str,
    context: str,
    row_unit: str,
    columns: list[dict[str, str]],
    created_at: str,
    sheet_order: int,
) -> dict[str, Any]:
    clean_name = str(name or "").strip()
    if not clean_name:
        raise ValueError("Sheet name is required.")
    clean_columns = normalize_columns(columns, require_questions=False)
    now = utc_now()
    return {
        "sheet_id": validate_sheet_id(sheet_id),
        "project_id": projects.validate_project_id(project_id),
        "name": clean_name,
        "context": str(context or "").strip(),
        "row_unit": str(row_unit or "").strip(),
        "columns": clean_columns,
        "sheet_order": int(sheet_order),
        "created_at": created_at or now,
        "updated_at": now,
        "locked_at": "",
        "locked_by_job_id": "",
        "locked_reason": "",
    }


def normalize_columns(columns: list[dict[str, Any]], *, require_questions: bool = True) -> list[dict[str, str]]:
    if not columns:
        raise ValueError("Add at least one structured column.")
    normalized: list[dict[str, str]] = []
    seen: set[str] = set()
    for index, column in enumerate(columns, start=1):
        column_name = str(column.get("column_name") or column.get("name") or "").strip()
        question = str(column.get("question") or "").strip()
        rules = str(column.get("rules") or "").strip()
        if not column_name:
            raise ValueError(f"Column {index} is missing a column name.")
        if "\t" in column_name or "\n" in column_name or "\r" in column_name:
            raise ValueError(f"Column `{column_name}` cannot contain tabs or line breaks.")
        if column_name in seen:
            raise ValueError(f"Duplicate column name: {column_name}")
        if require_questions and not question:
            raise ValueError(f"Column `{column_name}` is missing a question.")
        seen.add(column_name)
        normalized.append({"column_name": column_name, "question": question, "rules": rules})
    return normalized


def read_sheet_file(path: Path) -> dict[str, Any]:
    payload = jobs.read_json(path)
    payload["sheet_id"] = validate_sheet_id(str(payload.get("sheet_id") or path.parent.name))
    payload["project_id"] = projects.validate_project_id(str(payload.get("project_id") or path.parents[2].name))
    payload["columns"] = normalize_columns(payload.get("columns") or [], require_questions=False)
    payload.setdefault("context", "")
    payload.setdefault("row_unit", "")
    payload.setdefault("sheet_order", 1_000_000)
    payload.setdefault("created_at", "")
    payload.setdefault("updated_at", "")
    payload.setdefault("locked_at", "")
    payload.setdefault("locked_by_job_id", "")
    payload.setdefault("locked_reason", "")
    payload["is_locked"] = sheet_is_locked(payload)
    try:
        payload["compiled_prompt"] = compile_sheet_prompt(payload)
        payload["schema_ready"] = True
        payload["schema_error"] = ""
    except ValueError as exc:
        payload["compiled_prompt"] = ""
        payload["schema_ready"] = False
        payload["schema_error"] = str(exc)
    return payload


def sheet_is_locked(sheet: dict[str, Any]) -> bool:
    return bool(str(sheet.get("locked_at") or "").strip())


def ensure_sheet_editable(sheet: dict[str, Any]) -> None:
    if sheet_is_locked(sheet):
        raise ValueError("This sheet is locked for reproducibility. Duplicate it to edit the schema.")


def ensure_sheet_data_files(root: Path) -> None:
    for filename in [ROWS_JSONL_FILENAME, ERRORS_JSONL_FILENAME]:
        path = root / filename
        if not path.exists():
            path.write_text("", encoding="utf-8")


def extract_block_sections(block: str) -> dict[str, str]:
    matches = list(_SECTION_RE.finditer(block))
    if not matches:
        raise ValueError("Column block is missing required headings.")
    sections: dict[str, str] = {}
    for index, match in enumerate(matches):
        heading = match.group(1)
        inline_value = match.group(2) or ""
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(block)
        body = block[start:end]
        value = "\n".join(part for part in [inline_value.strip(), body.strip()] if part).strip()
        sections[heading] = value
    return sections


def strip_surrounding_code_fence(text: str) -> str:
    lines = text.strip(" \r\n").splitlines()
    if len(lines) >= 2 and lines[0].strip().startswith("```") and lines[-1].strip() == "```":
        return "\n".join(lines[1:-1]).strip(" \r\n")
    return text


def structured_row_record(
    *,
    project_id: str,
    sheet_id: str,
    job_id: str,
    metadata: dict[str, Any],
    extracted_at: str,
    model: str,
    cells: dict[str, str],
    row_index: int,
) -> dict[str, Any]:
    return {
        "row_id": uuid.uuid4().hex,
        "sheet_id": sheet_id,
        "project_id": project_id,
        "job_id": job_id,
        "source_pdf": str(metadata.get("original_filename") or ""),
        "source_relative_path": str(metadata.get("source_relative_path") or ""),
        "extracted_at": extracted_at,
        "parse_date": extracted_at,
        "model": model,
        "parse_status": "parsed",
        "parse_error": "",
        "row_index": row_index,
        "cells": cells,
    }


def structured_error_record(
    *,
    project_id: str,
    sheet_id: str,
    job_id: str,
    metadata: dict[str, Any],
    extracted_at: str,
    model: str,
    parse_error: str,
    raw_output_file: Path,
    raw_output: str,
) -> dict[str, Any]:
    return {
        "error_id": uuid.uuid4().hex,
        "sheet_id": sheet_id,
        "project_id": project_id,
        "job_id": job_id,
        "source_pdf": str(metadata.get("original_filename") or ""),
        "source_relative_path": str(metadata.get("source_relative_path") or ""),
        "extracted_at": extracted_at,
        "parse_date": extracted_at,
        "model": model,
        "parse_status": "failed",
        "parse_error": parse_error,
        "raw_output_file": str(raw_output_file),
        "raw_output_preview": raw_output[:1000],
        "cells": {},
    }


def replace_job_sheet_records(
    project_id: str,
    sheet_id: str,
    job_id: str,
    rows: list[dict[str, Any]],
    errors: list[dict[str, Any]],
) -> None:
    root = sheet_root(project_id, sheet_id)
    ensure_sheet_data_files(root)
    existing_rows = [row for row in read_jsonl(root / ROWS_JSONL_FILENAME) if str(row.get("job_id") or "") != job_id]
    existing_errors = [error for error in read_jsonl(root / ERRORS_JSONL_FILENAME) if str(error.get("job_id") or "") != job_id]
    write_jsonl(root / ROWS_JSONL_FILENAME, existing_rows + rows)
    write_jsonl(root / ERRORS_JSONL_FILENAME, existing_errors + errors)


def read_sheet_rows(project_id: str, sheet_id: str) -> list[dict[str, Any]]:
    return read_jsonl(sheet_root(project_id, sheet_id) / ROWS_JSONL_FILENAME)


def read_sheet_errors(project_id: str, sheet_id: str) -> list[dict[str, Any]]:
    return read_jsonl(sheet_root(project_id, sheet_id) / ERRORS_JSONL_FILENAME)


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    records: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            if not line.strip():
                continue
            records.append(json.loads(line))
    return records


def write_jsonl(path: Path, records: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    try:
        with tmp.open("w", encoding="utf-8") as handle:
            for record in records:
                handle.write(json.dumps(record, ensure_ascii=False) + "\n")
        tmp.replace(path)
    finally:
        try:
            tmp.unlink()
        except FileNotFoundError:
            pass


def row_search_haystack(record: dict[str, Any]) -> str:
    values = [
        record.get("source_pdf"),
        record.get("job_id"),
        record.get("parse_status"),
        record.get("parse_date"),
        record.get("parse_error"),
        record.get("job_status"),
        record.get("status_message"),
        record.get("raw_output_preview"),
    ]
    cells = record.get("cells") if isinstance(record.get("cells"), dict) else {}
    values.extend(cells.values())
    return " ".join(str(value or "") for value in values).lower()


def safe_worksheet_name(name: str, used_names: set[str]) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\]", " ", excel_safe_text(name or "Sheet")).strip() or "Sheet"
    base = cleaned[:31] or "Sheet"
    candidate = base
    suffix = 1
    while candidate in used_names:
        suffix_text = f" {suffix}"
        candidate = f"{base[:31 - len(suffix_text)]}{suffix_text}"
        suffix += 1
    used_names.add(candidate)
    return candidate


def excel_safe_text(value: Any) -> str:
    """Strip characters that are legal in JSON/text files but invalid in XLSX XML."""
    text = "" if value is None else str(value)
    return "".join(character for character in text if is_xml_compatible_character(character))


def is_xml_compatible_character(character: str) -> bool:
    codepoint = ord(character)
    return (
        codepoint in {0x09, 0x0A, 0x0D}
        or 0x20 <= codepoint <= 0xD7FF
        or 0xE000 <= codepoint <= 0xFFFD
        or 0x10000 <= codepoint <= 0x10FFFF
    )


def read_text_if_exists(path: Path) -> str:
    if not path.exists():
        return ""
    return path.read_text(encoding="utf-8")


def read_json_if_exists(path: Path, default: Any) -> Any:
    if not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8"))


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()
