from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


SUPPORTED_ATTACHMENT_SUFFIXES = {".pdf", ".xlsx", ".csv"}
MAX_SPREADSHEET_ROWS_PER_SHEET = 10_000
MAX_SPREADSHEET_CELLS_PER_FILE = 50_000
MAX_CELL_CHARACTERS = 4_000


@dataclass(frozen=True)
class SupplementaryTranscript:
    filename: str
    output_path: Path | None
    text: str
    warning: str = ""


def supported_attachment_suffix(filename: str) -> str:
    suffix = Path(str(filename or "")).suffix.lower()
    return suffix if suffix in SUPPORTED_ATTACHMENT_SUFFIXES else ""


def transcript_spreadsheet_sources(
    records: Iterable[dict[str, Any]],
    output_dir: Path,
) -> list[SupplementaryTranscript]:
    """Create reproducible text transcripts for CSV/XLSX supporting sources.

    The bounded conversion keeps one unusually large spreadsheet from turning a
    PDF extraction prompt into an unmanageable request. The source workbook is
    still retained unchanged in the job input directory.
    """

    transcripts: list[SupplementaryTranscript] = []
    for index, record in enumerate(records, start=1):
        filename = str(record.get("filename") or f"supporting-file-{index}")
        source_path = Path(record["path"])
        try:
            text = spreadsheet_to_text(source_path, filename)
        except Exception as exc:
            transcripts.append(
                SupplementaryTranscript(
                    filename=filename,
                    output_path=None,
                    text="",
                    warning=f"Could not read supporting spreadsheet {filename}: {exc}",
                )
            )
            continue
        target = output_dir / f"{index:03d}-{safe_transcript_name(filename)}.txt"
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(text, encoding="utf-8")
        transcripts.append(SupplementaryTranscript(filename=filename, output_path=target, text=text))
    return transcripts


def spreadsheet_to_text(path: Path, display_name: str) -> str:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        sheet_text, truncated = csv_to_text(path)
        return format_transcript(display_name, [("CSV", sheet_text)], truncated)
    if suffix == ".xlsx":
        sheets, truncated = xlsx_to_text(path)
        return format_transcript(display_name, sheets, truncated)
    raise ValueError(f"Unsupported supporting spreadsheet type: {suffix or 'unknown'}")


def csv_to_text(path: Path) -> tuple[str, bool]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        sample = handle.read(4096)
        handle.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        return rows_to_text(csv.reader(handle, dialect), MAX_SPREADSHEET_CELLS_PER_FILE)


def xlsx_to_text(path: Path) -> tuple[list[tuple[str, str]], bool]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    remaining_cells = MAX_SPREADSHEET_CELLS_PER_FILE
    truncated = False
    sheets: list[tuple[str, str]] = []
    try:
        for worksheet in workbook.worksheets:
            if remaining_cells <= 0:
                truncated = True
                break
            rows = worksheet.iter_rows(values_only=True)
            text, rows_truncated, cells_used = rows_to_text_with_usage(
                rows,
                max_cells=remaining_cells,
                max_rows=MAX_SPREADSHEET_ROWS_PER_SHEET,
            )
            remaining_cells -= cells_used
            truncated = truncated or rows_truncated
            sheets.append((str(worksheet.title or "Worksheet"), text))
    finally:
        workbook.close()
    return sheets, truncated


def rows_to_text(rows: Iterable[Iterable[Any]], max_cells: int) -> tuple[str, bool]:
    text, truncated, _cells_used = rows_to_text_with_usage(rows, max_cells=max_cells, max_rows=MAX_SPREADSHEET_ROWS_PER_SHEET)
    return text, truncated


def rows_to_text_with_usage(
    rows: Iterable[Iterable[Any]],
    *,
    max_cells: int,
    max_rows: int,
) -> tuple[str, bool, int]:
    output: list[str] = []
    cells_used = 0
    rows_seen = 0
    truncated = False
    for raw_row in rows:
        if rows_seen >= max_rows or cells_used >= max_cells:
            truncated = True
            break
        values = [normalize_cell_value(value) for value in raw_row]
        if not any(values):
            continue
        available = max_cells - cells_used
        if len(values) > available:
            values = values[:available]
            truncated = True
        output.append("\t".join(values))
        cells_used += len(values)
        rows_seen += 1
        if truncated:
            break
    return "\n".join(output), truncated, cells_used


def normalize_cell_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        text = value.isoformat()
    else:
        text = str(value)
    text = text.replace("\r\n", " ").replace("\r", " ").replace("\n", " ").replace("\t", " ")
    if len(text) > MAX_CELL_CHARACTERS:
        return text[:MAX_CELL_CHARACTERS] + " [cell truncated]"
    return text


def format_transcript(display_name: str, sheets: list[tuple[str, str]], truncated: bool) -> str:
    parts = [f"===== Supporting spreadsheet: {display_name} ====="]
    for sheet_name, text in sheets:
        parts.append(f"===== Worksheet: {sheet_name} =====")
        parts.append(text or "[No non-empty cells found]")
    if truncated:
        parts.append(
            "[Transcript truncated after the configured row/cell limit. The original supporting file remains in the job input directory.]"
        )
    return "\n".join(parts).strip() + "\n"


def safe_transcript_name(filename: str) -> str:
    base = Path(filename).stem or "supporting-spreadsheet"
    cleaned = "".join(character if character.isalnum() or character in {"-", "_"} else "-" for character in base)
    return cleaned.strip("-") or "supporting-spreadsheet"
