from __future__ import annotations

import csv
import json
import logging
import re
import shutil
import threading
import time
import uuid
from collections import deque
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook

from app import config
from app.models import JobSettings
from app.services import jobs, projects
from app.services.openai_rq import run_openai_rq
from app.services.rq_llm import run_rq_llm
from app.services.rq_prompt import build_prompt_transcript, read_prompt_file

logger = logging.getLogger(__name__)

TEXT_EXTRACTION_TYPE = "text"
SOURCE_FILES_DIRNAME = "source_files"
ROWS_JSONL_FILENAME = "rows.jsonl"
SOURCE_MANIFEST_FILENAME = "source.json"
TEXT_EXPORT_SUFFIX = "text_extraction_export.xlsx"
TEXT_JOB_OUTPUT = "output.md"
TEXT_SHEETS_DIRNAME = "text_sheets"
TEXT_SHEET_FILENAME = "sheet.json"
TEXT_WORKBOOK_FILENAME = "text_workbook.json"
VISIBLE_RESULT_PREVIEW_CHARS = 220
CSV_FIELD_LIMIT = 1024 * 1024 * 32
MAX_LOCAL_TEXT_CONCURRENT_JOBS = 4
_TEXT_SHEET_ID_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]{0,127}$")
_LOCAL_TEXT_INFERENCE_SEMAPHORE = threading.BoundedSemaphore(MAX_LOCAL_TEXT_CONCURRENT_JOBS)

try:
    csv.field_size_limit(CSV_FIELD_LIMIT)
except OverflowError:
    pass


@dataclass(frozen=True)
class TextQueuedJob:
    project_id: str
    job_id: str


class TextJobQueue:
    def __init__(self) -> None:
        self._pending: deque[TextQueuedJob] = deque()
        self._queued_or_running: set[str] = set()
        self._current: dict[str, TextQueuedJob] = {}
        self._paused = False
        self._condition = threading.Condition()
        worker_count = max(1, config.MAX_OPENAI_CONCURRENT_REQUESTS)
        self._workers = [
            threading.Thread(
                target=self._run,
                name=f"text-extraction-job-queue-{index + 1}",
                daemon=True,
            )
            for index in range(worker_count)
        ]
        for worker in self._workers:
            worker.start()

    def enqueue(self, project_id: str, job_id: str, *, front: bool = False) -> bool:
        key = queue_key(project_id, job_id)
        with self._condition:
            if key in self._queued_or_running:
                return False
            self._queued_or_running.add(key)
            queued = TextQueuedJob(project_id=project_id, job_id=job_id)
            if front:
                self._pending.appendleft(queued)
            else:
                self._pending.append(queued)
            self._condition.notify_all()
        jobs.update_status(
            job_id,
            status="queued",
            stage="queued",
            message="Queued",
            progress=0.0,
            error=None,
            event={"event": "queued"},
            project_id=project_id,
        )
        return True

    def active_job_ids(self, *, project_id: str | None = None) -> set[str]:
        with self._condition:
            return {
                job_id_from_key(key)
                for key in self._queued_or_running
                if project_id is None or key.startswith(f"{project_id}:")
            }

    def pause(self, *, project_id: str | None = None) -> dict[str, Any]:
        # Text pause/resume stops future row starts. A row already inside a model
        # call is allowed to finish so pause does not turn active work into a
        # failed retry.
        with self._condition:
            self._paused = True
            self._condition.notify_all()
        return self.status(project_id=project_id)

    def resume(self, *, project_id: str | None = None) -> dict[str, Any]:
        with self._condition:
            self._paused = False
            self._condition.notify_all()
        return self.status(project_id=project_id)

    def status(self, *, project_id: str | None = None) -> dict[str, Any]:
        with self._condition:
            current = [
                queued.job_id
                for queued in self._current.values()
                if project_id is None or queued.project_id == project_id
            ]
            pending = [
                queued.job_id
                for queued in self._pending
                if project_id is None or queued.project_id == project_id
            ]
        return {
            "paused": self._paused,
            "current_job_ids": current,
            "pending_job_ids": pending,
            "pending_count": len(pending),
            "running_count": len(current),
            "max_text_workers": len(self._workers),
        }

    def enqueue_existing_queued_jobs(self) -> int:
        self.mark_stale_running_jobs_queued()
        restored = 0
        for project_id in known_text_project_ids():
            records = sorted(
                text_job_records(project_id),
                key=lambda item: text_row_order(item.get("metadata") or {}),
            )
            for record in records:
                status = record.get("status") or {}
                if normalize_status_key(status.get("status")) != "queued":
                    continue
                if self.enqueue(project_id, str(record["job_id"])):
                    restored += 1
        if restored:
            logger.info("Restored %s queued text extraction jobs from disk", restored)
        return restored

    def mark_stale_running_jobs_queued(self, *, project_id: str | None = None) -> list[dict[str, Any]]:
        active_ids = self.active_job_ids(project_id=project_id)
        marked: list[dict[str, Any]] = []
        project_ids = [project_id] if project_id else known_text_project_ids()
        for active_project_id in project_ids:
            for record in text_job_records(active_project_id):
                status = record.get("status") or {}
                job_id = str(record["job_id"])
                if normalize_status_key(status.get("status")) != "running" or job_id in active_ids:
                    continue
                message = "Text extraction was interrupted while the app was restarting and has been returned to the queue."
                jobs.update_status(
                    job_id,
                    status="queued",
                    stage="queued",
                    message=message,
                    progress=0.0,
                    error=None,
                    event={"event": "stale_text_job_requeued"},
                    project_id=active_project_id,
                )
                marked.append(record)
        return marked

    def _run(self) -> None:
        while True:
            queued = self._next_job()
            key = queue_key(queued.project_id, queued.job_id)
            try:
                logger.info("Dequeued text extraction job %s", queued.job_id)
                run_text_job(queued.project_id, queued.job_id)
            except Exception:
                logger.exception("Queued text extraction job %s crashed outside job handling", queued.job_id)
                jobs.update_status(
                    queued.job_id,
                    status="failed",
                    stage="error",
                    message="Text extraction job crashed unexpectedly.",
                    progress=1.0,
                    error="Text extraction job crashed unexpectedly.",
                    event={"event": "error", "message": "Text extraction job crashed unexpectedly."},
                    project_id=queued.project_id,
                )
            finally:
                with self._condition:
                    if self._current.get(key) == queued:
                        self._current.pop(key, None)
                    if key not in {queue_key(item.project_id, item.job_id) for item in self._pending}:
                        self._queued_or_running.discard(key)
                    self._condition.notify_all()

    def _next_job(self) -> TextQueuedJob:
        with self._condition:
            while self._paused or not self._pending:
                self._condition.wait()
            queued = self._pending.popleft()
            self._current[queue_key(queued.project_id, queued.job_id)] = queued
            return queued


text_job_queue = TextJobQueue()


def get_text_workbook(project_id: str) -> dict[str, Any]:
    require_text_project(project_id)
    migrate_legacy_text_workbook(project_id)
    manifest = read_json_if_exists(text_workbook_path(project_id))
    source_id = str(manifest.get("source_id") or "")
    source = source_manifest(project_id, source_id) if source_id else None
    return {
        "project_id": project_id,
        "source": source,
        "sheets": list_text_sheets(project_id, migrate=False),
        "mapping_locked": bool(source),
        "created_at": str(manifest.get("created_at") or ""),
        "updated_at": str(manifest.get("updated_at") or ""),
    }


def create_text_source(
    *,
    project_id: str,
    upload_file: Any,
    column_mappings: list[dict[str, str]],
    study_id_column: str = "",
    initial_settings: JobSettings | None = None,
) -> dict[str, Any]:
    require_text_project(project_id)
    migrate_legacy_text_workbook(project_id)
    existing = read_json_if_exists(text_workbook_path(project_id))
    if existing.get("source_id"):
        raise ValueError("This project already has a spreadsheet and frozen column mapping.")

    original_filename = Path(str(upload_file.filename or "uploaded_spreadsheet")).name
    suffix = Path(original_filename).suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        if suffix == ".xls":
            raise ValueError("Legacy .xls files are not supported. Save the spreadsheet as .xlsx or CSV and upload again.")
        raise ValueError("Upload a .csv or .xlsx spreadsheet.")

    source_id = uuid.uuid4().hex
    source_dir = source_dir_for(project_id, source_id)
    source_dir.mkdir(parents=True, exist_ok=True)
    source_file = source_dir / safe_source_filename(original_filename)
    jobs.save_upload(upload_file.file, source_file)

    columns = read_spreadsheet_columns(source_file)
    mappings = normalize_column_mappings(column_mappings, columns)
    clean_study_id_column = normalize_optional_column(study_id_column, columns)
    rows_path = source_dir / ROWS_JSONL_FILENAME
    row_count = 0
    with rows_path.open("w", encoding="utf-8") as rows_handle:
        for source_row_index, row in enumerate(iter_spreadsheet_rows(source_file, columns)):
            if is_blank_row(row):
                continue
            row_display_number = source_row_index + 1
            row_payload = {
                "source_id": source_id,
                "source_filename": original_filename,
                "source_row_index": source_row_index,
                "row_display_number": row_display_number,
                "spreadsheet_row_number": source_row_index + 2,
                "record_id": record_id_for_row(row, clean_study_id_column, row_display_number),
                "study_id_column": clean_study_id_column,
                "selected_fields": selected_fields_for_row(row, mappings),
                "row": row,
            }
            rows_handle.write(json.dumps(row_payload, ensure_ascii=False) + "\n")
            row_count += 1

    now = utc_now()
    manifest = {
        "source_id": source_id,
        "source_filename": original_filename,
        "stored_source_file": str(source_file),
        "created_at": now,
        "columns": columns,
        "column_mappings": mappings,
        "study_id_column": clean_study_id_column,
        "row_count": row_count,
        "rows_jsonl": str(rows_path),
    }
    jobs.write_json(source_dir / SOURCE_MANIFEST_FILENAME, manifest)
    jobs.write_json(
        text_workbook_path(project_id),
        {
            "project_id": project_id,
            "source_id": source_id,
            "mapping_locked": True,
            "created_at": now,
            "updated_at": now,
        },
    )
    settings = initial_settings or JobSettings()
    sheet = create_text_sheet(
        project_id=project_id,
        name="Sheet 1",
        source_id=source_id,
        rq_model_preset=settings.rq_model_preset,
        rq_prompt_filename=settings.rq_prompt_filename,
        rq_system_prompt=settings.rq_system_prompt,
    )
    return {"source": manifest, "sheet": sheet, "sheets": list_text_sheets(project_id)}


def list_text_sheets(project_id: str, *, migrate: bool = True) -> list[dict[str, Any]]:
    require_text_project(project_id)
    if migrate:
        migrate_legacy_text_workbook(project_id)
    root = text_sheets_root(project_id)
    sheets: list[dict[str, Any]] = []
    for path in root.glob(f"*/{TEXT_SHEET_FILENAME}") if root.exists() else []:
        try:
            sheets.append(read_text_sheet_file(path))
        except Exception:
            logger.exception("Could not read text extraction sheet %s", path)
    sheets.sort(key=text_sheet_sort_key)
    return sheets


def get_text_sheet(project_id: str, sheet_id: str) -> dict[str, Any]:
    require_text_project(project_id)
    safe_sheet_id = validate_text_sheet_id(sheet_id)
    path = text_sheet_root(project_id, safe_sheet_id) / TEXT_SHEET_FILENAME
    if not path.exists():
        raise FileNotFoundError(f"Text extraction sheet not found: {safe_sheet_id}")
    return read_text_sheet_file(path)


def create_text_sheet(
    *,
    project_id: str,
    name: str = "",
    source_id: str = "",
    rq_model_preset: str = "qwen35_9b_8bit_reasoning",
    rq_prompt_filename: str = "",
    rq_system_prompt: str = "",
    insert_after_sheet_id: str = "",
) -> dict[str, Any]:
    require_text_project(project_id)
    workbook = read_json_if_exists(text_workbook_path(project_id))
    active_source_id = str(source_id or workbook.get("source_id") or "")
    if not active_source_id or not source_manifest(project_id, active_source_id):
        raise ValueError("Upload a spreadsheet before creating extraction sheets.")
    sheets = list_text_sheets(project_id, migrate=False)
    clean_name = unique_text_sheet_name(sheets, name or f"Sheet {len(sheets) + 1}")
    order = max([text_sheet_order(sheet) for sheet in sheets], default=-1) + 1
    if insert_after_sheet_id:
        existing = get_text_sheet(project_id, insert_after_sheet_id)
        order = text_sheet_order(existing) + 1
        shift_text_sheet_orders(project_id, start_order=order)
    now = utc_now()
    payload = {
        "sheet_id": unique_text_sheet_id(project_id, clean_name),
        "project_id": project_id,
        "source_id": active_source_id,
        "name": clean_name,
        "sheet_order": order,
        "rq_model_preset": str(rq_model_preset or "qwen35_9b_8bit_reasoning"),
        "rq_prompt_filename": str(rq_prompt_filename or ""),
        "rq_system_prompt": str(rq_system_prompt or ""),
        "rq_screening_model": "",
        "created_at": now,
        "updated_at": now,
        "locked_at": "",
        "locked_by_job_id": "",
    }
    root = text_sheet_root(project_id, str(payload["sheet_id"]))
    root.mkdir(parents=True, exist_ok=True)
    jobs.write_json(root / TEXT_SHEET_FILENAME, payload)
    return read_text_sheet_file(root / TEXT_SHEET_FILENAME)


def update_text_sheet(
    *,
    project_id: str,
    sheet_id: str,
    name: str,
    rq_model_preset: str,
    rq_prompt_filename: str,
    rq_system_prompt: str,
) -> dict[str, Any]:
    existing = get_text_sheet(project_id, sheet_id)
    ensure_text_sheet_editable(existing)
    clean_name = str(name or "").strip()
    if not clean_name:
        raise ValueError("Sheet name is required.")
    sibling_names = {
        str(sheet.get("name") or "")
        for sheet in list_text_sheets(project_id, migrate=False)
        if str(sheet.get("sheet_id")) != sheet_id
    }
    if clean_name in sibling_names:
        raise ValueError(f"A sheet named `{clean_name}` already exists.")
    existing.update(
        {
            "name": clean_name,
            "rq_model_preset": str(rq_model_preset or "qwen35_9b_8bit_reasoning"),
            "rq_prompt_filename": str(rq_prompt_filename or ""),
            "rq_system_prompt": str(rq_system_prompt or ""),
            "updated_at": utc_now(),
        }
    )
    write_text_sheet(existing)
    return get_text_sheet(project_id, sheet_id)


def duplicate_text_sheet(project_id: str, sheet_id: str) -> dict[str, Any]:
    existing = get_text_sheet(project_id, sheet_id)
    duplicate_name = next_text_duplicate_name(project_id, str(existing.get("name") or "Sheet"))
    return create_text_sheet(
        project_id=project_id,
        name=duplicate_name,
        source_id=str(existing.get("source_id") or ""),
        rq_model_preset=str(existing.get("rq_model_preset") or "qwen35_9b_8bit_reasoning"),
        rq_prompt_filename=str(existing.get("rq_prompt_filename") or ""),
        rq_system_prompt=str(existing.get("rq_system_prompt") or ""),
        insert_after_sheet_id=sheet_id,
    )


def create_text_jobs_for_sheet(*, project_id: str, sheet_id: str, settings: JobSettings) -> dict[str, Any]:
    sheet = get_text_sheet(project_id, sheet_id)
    ensure_text_sheet_editable(sheet)
    source = source_manifest(project_id, str(sheet.get("source_id") or ""))
    if not source:
        raise FileNotFoundError("The spreadsheet source for this sheet is missing.")
    if text_job_records(project_id, sheet_id=sheet_id):
        raise ValueError("This sheet already has row jobs. Duplicate it to run a different prompt or model.")

    rows = read_source_row_payloads(project_id, str(sheet["source_id"]))
    if not rows:
        raise ValueError("The uploaded spreadsheet has no nonblank data rows.")
    system_prompt, prompt_filename, prompt_source_path = resolve_system_prompt(settings)
    created_jobs: list[dict[str, Any]] = []
    for row_order, row_payload in enumerate(rows):
        selected_fields = row_payload.get("selected_fields") or []
        user_prompt = build_user_prompt(selected_fields)
        record_id = str(row_payload.get("record_id") or f"Row {row_payload.get('row_display_number') or row_order + 1}")
        job_id = jobs.new_job_id()
        root = jobs.create_job(job_id, record_id, settings, project_id=project_id)
        write_text_job_files(root=root, row_payload=row_payload, system_prompt=system_prompt, user_prompt=user_prompt)
        jobs.update_metadata(
            root,
            extraction_type=TEXT_EXTRACTION_TYPE,
            project_id=project_id,
            text_sheet_id=sheet_id,
            text_sheet_name=str(sheet.get("name") or sheet_id),
            source_id=str(sheet["source_id"]),
            source_filename=str(source.get("source_filename") or ""),
            stored_source_file=str(source.get("stored_source_file") or ""),
            original_filename=record_id,
            record_id=record_id,
            study_id_column=str(source.get("study_id_column") or ""),
            source_row_index=int(row_payload.get("source_row_index") or 0),
            row_display_number=int(row_payload.get("row_display_number") or row_order + 1),
            spreadsheet_row_number=int(row_payload.get("spreadsheet_row_number") or row_order + 2),
            project_row_order=row_order,
            sheet_row_order=row_order,
            input_preview=input_preview_for_fields(selected_fields, row_order + 1),
            selected_input_fields=source.get("column_mappings") or [],
            rq_prompt_source_path=prompt_source_path,
            rq_system_prompt_file=str(root / "system_prompt.txt"),
            rq_user_prompt_file=str(root / "user_prompt.txt"),
            **jobs.settings_metadata_updates(settings, rq_prompt_filename=prompt_filename),
        )
        created_jobs.append({"job_id": job_id, "record_id": record_id, "row_display_number": row_payload.get("row_display_number")})

    first_job_id = str(created_jobs[0]["job_id"])
    sheet.update(
        {
            "rq_model_preset": settings.rq_model_preset,
            "rq_prompt_filename": prompt_filename,
            "rq_system_prompt": system_prompt,
            "rq_screening_model": settings.rq_screening_model,
            "locked_at": utc_now(),
            "locked_by_job_id": first_job_id,
            "updated_at": utc_now(),
        }
    )
    write_text_sheet(sheet)
    for created in created_jobs:
        text_job_queue.enqueue(project_id, str(created["job_id"]))
    return {
        "sheet": get_text_sheet(project_id, sheet_id),
        "count": len(created_jobs),
        "job_id": first_job_id,
        **text_job_queue.status(project_id=project_id),
    }


def create_text_jobs_from_upload(
    *,
    project_id: str,
    upload_file: Any,
    settings: JobSettings,
    column_mappings: list[dict[str, str]],
    study_id_column: str = "",
) -> dict[str, Any]:
    source_result = create_text_source(
        project_id=project_id,
        upload_file=upload_file,
        column_mappings=column_mappings,
        study_id_column=study_id_column,
        initial_settings=settings,
    )
    run_result = create_text_jobs_for_sheet(
        project_id=project_id,
        sheet_id=str(source_result["sheet"]["sheet_id"]),
        settings=settings,
    )
    return {"source": source_result["source"], **run_result}


def run_text_job(project_id: str, job_id: str) -> None:
    root = jobs.job_dir(job_id, project_id)
    started = time.time()
    try:
        metadata = jobs.read_metadata(root)
        if metadata.get("extraction_type") != TEXT_EXTRACTION_TYPE:
            raise RuntimeError("Job is not a text extraction row job.")
        settings = jobs.load_job_settings(root)
        jobs.update_metadata(
            root,
            started_at=utc_now(),
            completed_at=None,
            duration_seconds=None,
            error_message=None,
        )
        jobs.update_status(
            job_id,
            status="running",
            stage="prompt",
            message="Preparing row prompt",
            progress=0.12,
            error=None,
            event={"event": "stage", "stage": "prompt"},
            project_id=project_id,
        )
        system_prompt_file = root / "system_prompt.txt"
        user_prompt_file = root / "user_prompt.txt"
        output_file = root / TEXT_JOB_OUTPUT
        if not system_prompt_file.exists() or not user_prompt_file.exists():
            raise RuntimeError("Stored row prompt files are missing.")

        if settings.rq_provider == "openai":
            jobs.update_status(
                job_id,
                status="running",
                stage="openai_running",
                message="OpenAI inference running",
                progress=0.36,
                event={"event": "openai_running"},
                project_id=project_id,
            )
            run_openai_rq(
                job_id=job_id,
                model=settings.rq_screening_model,
                system_prompt_file=system_prompt_file,
                user_prompt_file=user_prompt_file,
                output_file=output_file,
                max_tokens=settings.rq_max_tokens,
                enable_reasoning=settings.rq_enable_thinking,
                reasoning_effort=settings.openai_reasoning_effort,
                api_key=settings.openai_api_key,
                on_event=lambda event: handle_text_llm_event(job_id, event, provider="openai", project_id=project_id),
            )
        else:
            jobs.update_status(
                job_id,
                status="running",
                stage="rq_model",
                message="Loading extraction model",
                progress=0.28,
                event={"event": "stage", "stage": "rq_model"},
                project_id=project_id,
            )
            # The queue is sized for OpenAI throughput; retain the previous local
            # inference cap so increasing network concurrency does not overload MLX.
            with _LOCAL_TEXT_INFERENCE_SEMAPHORE:
                run_rq_llm(
                    job_id=job_id,
                    model=settings.rq_screening_model,
                    system_prompt_file=system_prompt_file,
                    user_prompt_file=user_prompt_file,
                    output_file=output_file,
                    max_tokens=settings.rq_max_tokens,
                    thinking_budget=settings.rq_thinking_budget,
                    temperature=settings.rq_temperature,
                    top_p=settings.rq_top_p,
                    top_k=settings.rq_top_k,
                    min_p=settings.rq_min_p,
                    presence_penalty=settings.rq_presence_penalty,
                    repetition_penalty=settings.rq_repetition_penalty,
                    enable_thinking=settings.rq_enable_thinking,
                    verbose=settings.rq_local_inference_verbose,
                    on_event=lambda event: handle_text_llm_event(job_id, event, provider="local", project_id=project_id),
                )

        completed_at = utc_now()
        duration_seconds = round(time.time() - started, 2)
        copy_output_for_compatibility(root)
        jobs.update_metadata(
            root,
            completed_at=completed_at,
            duration_seconds=duration_seconds,
            error_message=None,
            text_output_file=str(output_file),
            **jobs.settings_metadata_updates(settings, include_settings=False),
        )
        jobs.update_status(
            job_id,
            status="complete",
            stage="complete",
            message="Text extraction complete",
            progress=1.0,
            error=None,
            event={"event": "complete"},
            project_id=project_id,
        )
    except Exception as exc:
        logger.exception("Text extraction job %s failed", job_id)
        jobs.update_metadata(root, error_message=str(exc))
        jobs.update_status(
            job_id,
            status="failed",
            stage="error",
            message=str(exc),
            progress=1.0,
            error=str(exc),
            event={"event": "error", "message": str(exc)},
            project_id=project_id,
        )


def handle_text_llm_event(job_id: str, event: dict[str, Any], *, provider: str, project_id: str) -> None:
    name = str(event.get("event") or "")
    if name == "rq_model_loading":
        message = "Loading extraction model" if provider == "local" else "Starting OpenAI inference"
        progress = 0.36
    elif name == "rq_model_loaded":
        message = "Extraction model loaded" if provider == "local" else "OpenAI client ready"
        progress = 0.44
    elif name == "rq_generation_started":
        message = "Running row extraction" if provider == "local" else "OpenAI row extraction running"
        progress = 0.56
    elif name == "rq_generation_finished":
        message = "Row extraction complete"
        progress = 0.94
    elif name == "rq_generation_empty":
        message = "Model returned no final text"
        progress = 0.92
    elif name == "rq_generation_log":
        message = str(event.get("line") or "Model generating")
        progress = 0.72
    else:
        message = "Running row extraction"
        progress = 0.72
    jobs.update_status(
        job_id,
        status="running",
        stage="rq_screening" if provider == "local" else "openai_running",
        message=message,
        progress=progress,
        event=event,
        project_id=project_id,
    )


def list_text_jobs(
    *,
    project_id: str,
    sheet_id: str = "",
    offset: int = 0,
    limit: int = 100,
    status: str = "all",
    search: str = "",
) -> dict[str, Any]:
    require_text_project(project_id)
    workbook = get_text_workbook(project_id)
    sheets = workbook.get("sheets") or []
    active_sheet_id = str(sheet_id or (sheets[0].get("sheet_id") if sheets else ""))
    if not active_sheet_id:
        return {"items": [], "total": 0, "offset": 0, "limit": max(1, min(500, int(limit or 100))), "sheet_id": ""}
    sheet = get_text_sheet(project_id, active_sheet_id)
    text_job_queue.mark_stale_running_jobs_queued(project_id=project_id)
    normalized_status = normalize_status_key(status)
    query = str(search or "").strip().lower()
    source_id = str(sheet.get("source_id") or "")
    source = source_manifest(project_id, source_id)
    mappings = source.get("column_mappings") or []
    records_by_row = {
        int((record.get("metadata") or {}).get("source_row_index") or 0): record
        for record in text_job_records(project_id, sheet_id=active_sheet_id)
    }
    filtered: list[dict[str, Any]] = []
    for row_payload in read_source_row_payloads(project_id, source_id):
        row_index = int(row_payload.get("source_row_index") or 0)
        record = records_by_row.get(row_index)
        item = serialize_text_sheet_row(
            project_id=project_id,
            row_payload=row_payload,
            mappings=mappings,
            record=record,
        )
        status_key = str(item.get("status") or "not_run")
        if normalized_status != "all" and status_key != normalized_status:
            continue
        if query and query not in text_sheet_row_search_haystack(item):
            continue
        filtered.append(item)

    safe_offset = max(0, int(offset or 0))
    safe_limit = max(1, min(500, int(limit or 100)))
    window = filtered[safe_offset : safe_offset + safe_limit]
    return {
        "items": window,
        "total": len(filtered),
        "offset": safe_offset,
        "limit": safe_limit,
        "sheet_id": active_sheet_id,
    }


def text_job_counts(project_id: str, *, sheet_id: str = "") -> dict[str, int]:
    require_text_project(project_id)
    workbook = get_text_workbook(project_id)
    sheets = workbook.get("sheets") or []
    active_sheet_id = str(sheet_id or (sheets[0].get("sheet_id") if sheets else ""))
    text_job_queue.mark_stale_running_jobs_queued(project_id=project_id)
    counts = {"total": 0, "not_run": 0, "queued": 0, "running": 0, "completed": 0, "failed": 0}
    if not active_sheet_id:
        return counts
    sheet = get_text_sheet(project_id, active_sheet_id)
    source = source_manifest(project_id, str(sheet.get("source_id") or ""))
    counts["total"] = int(source.get("row_count") or 0)
    records = text_job_records(project_id, sheet_id=active_sheet_id)
    for record in records:
        key = normalize_status_key((record.get("status") or {}).get("status"))
        if key in counts and key != "total":
            counts[key] += 1
    counts["not_run"] = max(0, counts["total"] - len(records))
    return counts


def retry_text_job(project_id: str, job_id: str, *, sheet_id: str = "") -> dict[str, Any]:
    require_text_project(project_id)
    root = jobs.job_dir(job_id, project_id)
    if not root.exists():
        raise FileNotFoundError("Text row job not found.")
    metadata = jobs.read_metadata(root)
    if metadata.get("extraction_type") != TEXT_EXTRACTION_TYPE:
        raise ValueError("Job is not a text extraction row job.")
    if sheet_id and str(metadata.get("text_sheet_id") or "") != sheet_id:
        raise ValueError("Text row job does not belong to the selected sheet.")
    status = jobs.read_status(job_id, project_id=project_id)
    if normalize_status_key(status.get("status")) == "running" or job_id in text_job_queue.active_job_ids(project_id=project_id):
        raise ValueError("Row is already queued or running.")
    if normalize_status_key(status.get("status")) != "failed":
        raise ValueError("Only failed rows can be retried.")
    mark_text_job_for_retry(root, previous_status=status)
    text_job_queue.enqueue(project_id, job_id, front=True)
    return serialize_text_job(text_job_record(project_id, job_id), project_id=project_id)


def retry_failed_text_jobs(project_id: str, *, sheet_id: str = "") -> dict[str, Any]:
    require_text_project(project_id)
    requeued: list[dict[str, Any]] = []
    skipped: list[dict[str, str]] = []
    records = sorted(
        text_job_records(project_id, sheet_id=sheet_id),
        key=lambda item: text_row_order(item.get("metadata") or {}),
    )
    for record in records:
        status = record.get("status") or {}
        if normalize_status_key(status.get("status")) != "failed":
            continue
        job_id = str(record["job_id"])
        try:
            root = jobs.job_dir(job_id, project_id)
            jobs.load_original_job_settings(root)
            mark_text_job_for_retry(root, previous_status=status)
            if text_job_queue.enqueue(project_id, job_id):
                requeued.append(serialize_text_job(text_job_record(project_id, job_id), project_id=project_id))
        except Exception as exc:
            skipped.append({"job_id": job_id, "error": str(exc)})
    return {"requeued": len(requeued), "jobs": requeued, "skipped": skipped, **text_job_queue.status(project_id=project_id)}


def read_text_job_output(project_id: str, job_id: str, *, sheet_id: str = "") -> dict[str, Any]:
    require_text_project(project_id)
    root = jobs.job_dir(job_id, project_id)
    if not root.exists():
        raise FileNotFoundError("Text row job not found.")
    metadata = jobs.read_metadata(root)
    if metadata.get("extraction_type") != TEXT_EXTRACTION_TYPE:
        raise ValueError("Job is not a text extraction row job.")
    if sheet_id and str(metadata.get("text_sheet_id") or "") != sheet_id:
        raise ValueError("Text row job does not belong to the selected sheet.")
    return {
        "status": jobs.read_status(job_id, project_id=project_id),
        "metadata": metadata,
        "row_input": read_json_if_exists(root / "row_input.json"),
        "user_prompt": read_text_if_exists(root / "user_prompt.txt"),
        "system_prompt": read_text_if_exists(root / "system_prompt.txt"),
        "output": read_text_if_exists(root / TEXT_JOB_OUTPUT),
        "job_dir": str(root),
    }


def export_text_results(project_id: str) -> Path:
    project = require_text_project(project_id)
    workbook_state = get_text_workbook(project_id)
    sheets = workbook_state.get("sheets") or []
    if not sheets:
        raise ValueError("Create at least one extraction sheet before exporting.")
    cerebro_columns = [
        "cerebro_extracted_at",
        "cerebro_processing_time_seconds",
        "cerebro_model_used",
        "cerebro_prompt_used",
        "cerebro_status",
        "cerebro_error",
        "cerebro_output",
    ]
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_names: set[str] = set()
    for sheet in sheets:
        source_id = str(sheet.get("source_id") or "")
        source = source_manifest(project_id, source_id)
        original_columns = [str(column) for column in source.get("columns") or []]
        worksheet = workbook.create_sheet(safe_text_worksheet_name(str(sheet.get("name") or "Sheet"), used_names))
        worksheet.append(original_columns + cerebro_columns)
        records_by_row = {
            int((record.get("metadata") or {}).get("source_row_index") or 0): record
            for record in text_job_records(project_id, sheet_id=str(sheet.get("sheet_id") or ""))
        }
        for row_payload in read_source_row_payloads(project_id, source_id):
            row_index = int(row_payload.get("source_row_index") or 0)
            original_row = row_payload.get("row") or {}
            record = records_by_row.get(row_index)
            metadata = (record or {}).get("metadata") or {}
            status = (record or {}).get("status") or {}
            status_key = normalize_status_key(status.get("status")) if record else "not_run"
            output = ""
            if record and status_key == "completed":
                output = read_text_if_exists(jobs.job_dir(str(record["job_id"]), project_id) / TEXT_JOB_OUTPUT)
            worksheet.append(
                [original_row.get(column, "") for column in original_columns]
                + [
                    str(metadata.get("completed_at") or ""),
                    metadata.get("duration_seconds"),
                    str(metadata.get("rq_screening_model") or sheet.get("rq_screening_model") or ""),
                    str(metadata.get("rq_prompt_filename") or sheet.get("rq_prompt_filename") or ""),
                    status_label(status_key),
                    str(status.get("error") or metadata.get("error_message") or ""),
                    output,
                ]
            )
        for column in worksheet.columns:
            letter = column[0].column_letter
            worksheet.column_dimensions[letter].width = min(max(len(str(column[0].value or "")) + 4, 14), 60)
    path = projects.project_root(project_id) / f"cerebro_{projects.sanitize_slug(str(project.get('name') or project_id))}_{TEXT_EXPORT_SUFFIX}"
    workbook.save(path)
    return path


def require_text_project(project_id: str) -> dict[str, Any]:
    project = projects.get_project(project_id)
    if projects.normalize_extraction_type(project.get("extraction_type")) != TEXT_EXTRACTION_TYPE:
        raise ValueError("Project is not a text extraction project.")
    return project


def known_text_project_ids() -> list[str]:
    return [
        str(project["project_id"])
        for project in projects.list_projects()
        if projects.normalize_extraction_type(project.get("extraction_type")) == TEXT_EXTRACTION_TYPE
    ]


def migrate_legacy_text_workbook(project_id: str) -> dict[str, Any]:
    require_text_project(project_id)
    workbook_path = text_workbook_path(project_id)
    workbook = read_json_if_exists(workbook_path)
    all_records = text_job_records(project_id)
    sources = list_source_manifests(project_id)
    if not workbook and not sources and not all_records:
        return {}

    if not workbook:
        preferred_source_id = ""
        if all_records:
            preferred_source_id = str((all_records[0].get("metadata") or {}).get("source_id") or "")
        if not preferred_source_id and sources:
            preferred_source_id = str(sources[0].get("source_id") or "")
        now = utc_now()
        workbook = {
            "project_id": project_id,
            "source_id": preferred_source_id,
            "mapping_locked": bool(preferred_source_id),
            "created_at": str((sources[0] if sources else {}).get("created_at") or now),
            "updated_at": now,
            "migrated_from_legacy_jobs": bool(all_records),
        }
        jobs.write_json(workbook_path, workbook)

    sheets = list_text_sheets(project_id, migrate=False)
    for sheet in sheets:
        locked_job_id = str(sheet.get("locked_by_job_id") or "")
        if not locked_job_id:
            continue
        locked_root = jobs.job_dir(locked_job_id, project_id)
        if not locked_root.exists():
            continue
        locked_metadata = jobs.read_metadata(locked_root)
        locked_settings = locked_metadata.get("settings") or {}
        saved_preset = str(locked_metadata.get("rq_model_preset") or locked_settings.get("rq_model_preset") or "")
        if saved_preset and saved_preset != str(sheet.get("rq_model_preset") or ""):
            sheet["rq_model_preset"] = saved_preset
            sheet["updated_at"] = utc_now()
            write_text_sheet(sheet)
    orphan_records = [
        record for record in all_records if not str((record.get("metadata") or {}).get("text_sheet_id") or "")
    ]
    grouped: dict[tuple[str, str, str, str], list[dict[str, Any]]] = {}
    for record in orphan_records:
        metadata = record.get("metadata") or {}
        saved_settings = metadata.get("settings") or {}
        key = (
            str(metadata.get("source_id") or workbook.get("source_id") or ""),
            str(metadata.get("rq_prompt_filename") or ""),
            str(metadata.get("rq_model_preset") or saved_settings.get("rq_model_preset") or ""),
            str(metadata.get("rq_screening_model") or ""),
        )
        grouped.setdefault(key, []).append(record)

    for source_id, prompt_filename, model_preset, model_name in grouped:
        matching = next(
            (
                sheet
                for sheet in sheets
                if str(sheet.get("source_id") or "") == source_id
                and str(sheet.get("rq_prompt_filename") or "") == prompt_filename
                and str(sheet.get("rq_screening_model") or "") == model_name
            ),
            None,
        )
        records = sorted(grouped[(source_id, prompt_filename, model_preset, model_name)], key=lambda item: text_row_order(item.get("metadata") or {}))
        first_record = records[0]
        first_metadata = first_record.get("metadata") or {}
        if matching is None:
            raw_name = Path(prompt_filename).stem.replace("_", " ").strip() if prompt_filename else ""
            sheet_name = unique_text_sheet_name(sheets, raw_name or f"Sheet {len(sheets) + 1}")
            now = utc_now()
            payload = {
                "sheet_id": unique_text_sheet_id(project_id, sheet_name),
                "project_id": project_id,
                "source_id": source_id,
                "name": sheet_name,
                "sheet_order": len(sheets),
                "rq_model_preset": model_preset or "qwen35_9b_8bit_reasoning",
                "rq_prompt_filename": prompt_filename,
                "rq_system_prompt": read_text_if_exists(jobs.job_dir(str(first_record["job_id"]), project_id) / "system_prompt.txt").strip(),
                "rq_screening_model": model_name,
                "created_at": str(first_metadata.get("created_at") or now),
                "updated_at": now,
                "locked_at": str(first_metadata.get("created_at") or now),
                "locked_by_job_id": str(first_record["job_id"]),
                "migrated_from_legacy_jobs": True,
            }
            root = text_sheet_root(project_id, str(payload["sheet_id"]))
            root.mkdir(parents=True, exist_ok=True)
            jobs.write_json(root / TEXT_SHEET_FILENAME, payload)
            matching = read_text_sheet_file(root / TEXT_SHEET_FILENAME)
            sheets.append(matching)
        for record in records:
            root = jobs.job_dir(str(record["job_id"]), project_id)
            jobs.update_metadata(
                root,
                text_sheet_id=str(matching["sheet_id"]),
                text_sheet_name=str(matching.get("name") or matching["sheet_id"]),
                sheet_row_order=int((record.get("metadata") or {}).get("source_row_index") or 0),
            )
    return read_json_if_exists(workbook_path)


def text_workbook_path(project_id: str) -> Path:
    return projects.project_root(project_id) / TEXT_WORKBOOK_FILENAME


def text_sheets_root(project_id: str) -> Path:
    return projects.project_root(project_id) / TEXT_SHEETS_DIRNAME


def text_sheet_root(project_id: str, sheet_id: str) -> Path:
    return text_sheets_root(project_id) / validate_text_sheet_id(sheet_id)


def validate_text_sheet_id(sheet_id: str) -> str:
    candidate = str(sheet_id or "").strip()
    if not _TEXT_SHEET_ID_RE.fullmatch(candidate) or candidate in {".", ".."}:
        raise ValueError("Invalid text extraction sheet id.")
    return candidate


def unique_text_sheet_id(project_id: str, name: str) -> str:
    stem = projects.sanitize_slug(name)
    for _attempt in range(20):
        candidate = f"{stem}-{uuid.uuid4().hex[:8]}"
        if not text_sheet_root(project_id, candidate).exists():
            return candidate
    return uuid.uuid4().hex


def read_text_sheet_file(path: Path) -> dict[str, Any]:
    payload = jobs.read_json(path)
    payload["sheet_id"] = validate_text_sheet_id(str(payload.get("sheet_id") or path.parent.name))
    payload["project_id"] = projects.validate_project_id(str(payload.get("project_id") or path.parents[2].name))
    payload.setdefault("source_id", "")
    payload.setdefault("name", "Sheet")
    payload.setdefault("sheet_order", 1_000_000)
    payload.setdefault("rq_model_preset", "qwen35_9b_8bit_reasoning")
    payload.setdefault("rq_prompt_filename", "")
    payload.setdefault("rq_system_prompt", "")
    payload.setdefault("rq_screening_model", "")
    payload.setdefault("created_at", "")
    payload.setdefault("updated_at", "")
    payload.setdefault("locked_at", "")
    payload.setdefault("locked_by_job_id", "")
    payload["is_locked"] = bool(str(payload.get("locked_at") or ""))
    return payload


def write_text_sheet(sheet: dict[str, Any]) -> None:
    project_id = projects.validate_project_id(str(sheet.get("project_id") or ""))
    sheet_id = validate_text_sheet_id(str(sheet.get("sheet_id") or ""))
    payload = {key: value for key, value in sheet.items() if key != "is_locked"}
    jobs.write_json(text_sheet_root(project_id, sheet_id) / TEXT_SHEET_FILENAME, payload)


def ensure_text_sheet_editable(sheet: dict[str, Any]) -> None:
    if bool(str(sheet.get("locked_at") or "")):
        raise ValueError("This sheet has already been run. Duplicate it to change the prompt or model.")


def text_sheet_order(sheet: dict[str, Any]) -> int:
    try:
        return int(sheet.get("sheet_order"))
    except (TypeError, ValueError):
        return 1_000_000


def text_sheet_sort_key(sheet: dict[str, Any]) -> tuple[int, str, str]:
    return (text_sheet_order(sheet), str(sheet.get("created_at") or ""), str(sheet.get("sheet_id") or ""))


def shift_text_sheet_orders(project_id: str, *, start_order: int) -> None:
    for sheet in reversed(list_text_sheets(project_id, migrate=False)):
        if text_sheet_order(sheet) >= start_order:
            sheet["sheet_order"] = text_sheet_order(sheet) + 1
            write_text_sheet(sheet)


def unique_text_sheet_name(sheets: list[dict[str, Any]], requested: str) -> str:
    clean = str(requested or "Sheet").strip() or "Sheet"
    existing = {str(sheet.get("name") or "") for sheet in sheets}
    if clean not in existing:
        return clean
    base = re.sub(r"\s+\(\d+\)$", "", clean).strip() or "Sheet"
    index = 1
    while f"{base} ({index})" in existing:
        index += 1
    return f"{base} ({index})"


def next_text_duplicate_name(project_id: str, name: str) -> str:
    base = re.sub(r"\s+\(\d+\)$", "", str(name or "Sheet").strip()).strip() or "Sheet"
    existing = {str(sheet.get("name") or "") for sheet in list_text_sheets(project_id, migrate=False)}
    index = 1
    while f"{base} ({index})" in existing:
        index += 1
    return f"{base} ({index})"


def list_source_manifests(project_id: str) -> list[dict[str, Any]]:
    root = projects.project_root(project_id) / SOURCE_FILES_DIRNAME
    manifests: list[dict[str, Any]] = []
    for path in root.glob(f"*/{SOURCE_MANIFEST_FILENAME}") if root.exists() else []:
        payload = read_json_if_exists(path)
        if payload:
            manifests.append(payload)
    manifests.sort(key=lambda item: (str(item.get("created_at") or ""), str(item.get("source_id") or "")))
    return manifests


def source_manifest(project_id: str, source_id: str) -> dict[str, Any]:
    if not source_id:
        return {}
    return read_json_if_exists(source_dir_for(project_id, source_id) / SOURCE_MANIFEST_FILENAME)


def read_source_row_payloads(project_id: str, source_id: str) -> list[dict[str, Any]]:
    rows_path = source_dir_for(project_id, source_id) / ROWS_JSONL_FILENAME
    if not rows_path.exists():
        return []
    rows: list[dict[str, Any]] = []
    with rows_path.open("r", encoding="utf-8") as handle:
        for line in handle:
            if line.strip():
                rows.append(json.loads(line))
    return rows


def source_dir_for(project_id: str, source_id: str) -> Path:
    return projects.project_root(project_id) / SOURCE_FILES_DIRNAME / source_id


def safe_source_filename(filename: str) -> str:
    stem = projects.sanitize_slug(Path(filename).stem)
    suffix = Path(filename).suffix.lower()
    return f"{stem}{suffix or '.csv'}"


def read_spreadsheet_columns(path: Path) -> list[str]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            try:
                header = next(reader)
            except StopIteration:
                raise ValueError("Spreadsheet is empty.")
            return normalize_columns(header)
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet = workbook.active
            header = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
            if header is None:
                raise ValueError("Spreadsheet is empty.")
            return normalize_columns([stringify_cell(value) for value in header])
        finally:
            workbook.close()
    raise ValueError("Upload a .csv or .xlsx spreadsheet.")


def iter_spreadsheet_rows(path: Path, columns: list[str]) -> Iterable[dict[str, str]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            next(reader, None)
            for raw_values in reader:
                yield {column: stringify_cell(raw_values[index] if index < len(raw_values) else "") for index, column in enumerate(columns)}
        return
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        for values in worksheet.iter_rows(min_row=2, values_only=True):
            row_values = [stringify_cell(value) for value in values or []]
            row = {column: row_values[index] if index < len(row_values) else "" for index, column in enumerate(columns)}
            yield row
    finally:
        workbook.close()


def normalize_columns(header: list[Any]) -> list[str]:
    columns: list[str] = []
    seen: set[str] = set()
    for index, value in enumerate(header, start=1):
        column = stringify_cell(value).strip() or f"Column {index}"
        if column in seen:
            raise ValueError(f"Spreadsheet has a duplicate column name: {column}")
        columns.append(column)
        seen.add(column)
    if not columns:
        raise ValueError("Spreadsheet must include a header row.")
    return columns


def normalize_column_mappings(raw_mappings: list[dict[str, str]], columns: list[str]) -> list[dict[str, str]]:
    if not raw_mappings:
        raise ValueError("Add at least one column mapping.")
    available = set(columns)
    normalized: list[dict[str, str]] = []
    for index, mapping in enumerate(raw_mappings, start=1):
        column_name = str(mapping.get("column_name") or mapping.get("column") or "").strip()
        prompt_label = str(mapping.get("prompt_label") or mapping.get("label") or "").strip()
        if not column_name:
            raise ValueError(f"Column mapping {index} is missing a column name.")
        if not prompt_label:
            raise ValueError(f"Column mapping {index} is missing a prompt label.")
        if column_name not in available:
            raise ValueError(f"Spreadsheet column not found: {column_name}")
        normalized.append({"column_name": column_name, "prompt_label": prompt_label})
    return normalized


def normalize_optional_column(column_name: str, columns: list[str]) -> str:
    cleaned = str(column_name or "").strip()
    if not cleaned:
        return ""
    if cleaned not in set(columns):
        raise ValueError(f"Study ID column not found: {cleaned}")
    return cleaned


def selected_fields_for_row(row: dict[str, str], mappings: list[dict[str, str]]) -> list[dict[str, str]]:
    return [
        {
            "column_name": mapping["column_name"],
            "prompt_label": mapping["prompt_label"],
            "value": stringify_cell(row.get(mapping["column_name"], "")),
        }
        for mapping in mappings
    ]


def build_user_prompt(selected_fields: list[dict[str, str]]) -> str:
    return "\n\n".join(f"{field['prompt_label']}: {field.get('value', '')}" for field in selected_fields)


def input_preview_for_fields(selected_fields: list[dict[str, str]], row_display_number: int) -> str:
    for field in selected_fields:
        value = stringify_cell(field.get("value", "")).strip()
        if value:
            return truncate(value, 180)
    return f"Row {row_display_number}"


def record_id_for_row(row: dict[str, str], study_id_column: str, row_display_number: int) -> str:
    if study_id_column:
        value = stringify_cell(row.get(study_id_column, "")).strip()
        if value:
            return value
    return f"Row {row_display_number}"


def resolve_system_prompt(settings: JobSettings) -> tuple[str, str, str]:
    if settings.rq_system_prompt.strip():
        return settings.rq_system_prompt.strip(), settings.rq_prompt_filename, ""
    prompt_record = read_prompt_file(settings.rq_prompt_filename)
    return prompt_record["system_prompt"].strip(), prompt_record["filename"], prompt_record["path"]


def write_text_job_files(root: Path, row_payload: dict[str, Any], system_prompt: str, user_prompt: str) -> None:
    jobs.write_json(root / "row_input.json", row_payload)
    (root / "system_prompt.txt").write_text(system_prompt + "\n", encoding="utf-8")
    (root / "user_prompt.txt").write_text(user_prompt + "\n", encoding="utf-8")
    (root / "outputs" / "rq_system_prompt.txt").write_text(system_prompt + "\n", encoding="utf-8")
    (root / "outputs" / "rq_user_prompt.txt").write_text(user_prompt + "\n", encoding="utf-8")
    (root / "outputs" / "rq_prompt.txt").write_text(build_prompt_transcript(system_prompt, user_prompt), encoding="utf-8")


def copy_output_for_compatibility(root: Path) -> None:
    output = root / TEXT_JOB_OUTPUT
    if not output.exists():
        return
    compatibility_output = root / "outputs" / "rq_screening_output.md"
    compatibility_output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(output, compatibility_output)


def mark_text_job_for_retry(root: Path, *, previous_status: dict[str, Any]) -> None:
    for path in [root / TEXT_JOB_OUTPUT, root / "outputs" / "rq_screening_output.md"]:
        try:
            path.unlink()
        except FileNotFoundError:
            pass
    jobs.update_metadata(
        root,
        retry_failed_at=utc_now(),
        retry_failed_from_status=previous_status.get("status"),
        retry_failed_from_stage=previous_status.get("stage"),
        completed_at=None,
        duration_seconds=None,
        error_message=None,
    )


def text_job_records(project_id: str, *, sheet_id: str = "") -> list[dict[str, Any]]:
    records = [
        record
        for record in jobs.list_jobs(limit=0, project_id=project_id)
        if (record.get("metadata") or {}).get("extraction_type") == TEXT_EXTRACTION_TYPE
    ]
    if sheet_id:
        records = [
            record
            for record in records
            if str((record.get("metadata") or {}).get("text_sheet_id") or "") == sheet_id
        ]
    return records


def text_job_record(project_id: str, job_id: str) -> dict[str, Any]:
    for record in text_job_records(project_id):
        if str(record["job_id"]) == job_id:
            return record
    raise FileNotFoundError("Text row job not found.")


def serialize_text_job(record: dict[str, Any], *, project_id: str) -> dict[str, Any]:
    metadata = record.get("metadata") or {}
    status = record.get("status") or {}
    status_key = normalize_status_key(status.get("status"))
    root = jobs.job_dir(str(record["job_id"]), project_id)
    output = read_text_if_exists(root / TEXT_JOB_OUTPUT) if status_key == "completed" else ""
    return {
        "job_id": str(record["job_id"]),
        "project_id": project_id,
        "row": int(metadata.get("row_display_number") or 0),
        "row_order": text_row_order(metadata),
        "status": status_key,
        "status_label": status_label(status_key),
        "stage": str(status.get("stage") or ""),
        "message": str(status.get("message") or ""),
        "error": str(status.get("error") or metadata.get("error_message") or ""),
        "record_id": str(metadata.get("record_id") or f"Row {metadata.get('row_display_number') or ''}"),
        "input_preview": str(metadata.get("input_preview") or ""),
        "model": str(metadata.get("rq_screening_model") or ""),
        "prompt": str(metadata.get("rq_prompt_filename") or ""),
        "started_at": str(metadata.get("started_at") or ""),
        "completed_at": str(metadata.get("completed_at") or ""),
        "duration_seconds": metadata.get("duration_seconds"),
        "result_preview": truncate(output.strip(), VISIBLE_RESULT_PREVIEW_CHARS),
    }


def serialize_text_sheet_row(
    *,
    project_id: str,
    row_payload: dict[str, Any],
    mappings: list[dict[str, Any]],
    record: dict[str, Any] | None,
) -> dict[str, Any]:
    row = row_payload.get("row") or {}
    mapped_cells = {
        str(mapping.get("column_name") or ""): stringify_cell(row.get(str(mapping.get("column_name") or ""), ""))
        for mapping in mappings
    }
    if record is not None:
        item = serialize_text_job(record, project_id=project_id)
        item["mapped_cells"] = mapped_cells
        item["source_row_index"] = int(row_payload.get("source_row_index") or 0)
        return item
    row_display_number = int(row_payload.get("row_display_number") or 0)
    return {
        "job_id": "",
        "project_id": project_id,
        "row": row_display_number,
        "row_order": int(row_payload.get("source_row_index") or 0),
        "source_row_index": int(row_payload.get("source_row_index") or 0),
        "status": "not_run",
        "status_label": "Not run",
        "stage": "",
        "message": "",
        "error": "",
        "record_id": str(row_payload.get("record_id") or f"Row {row_display_number}"),
        "input_preview": input_preview_for_fields(row_payload.get("selected_fields") or [], row_display_number),
        "mapped_cells": mapped_cells,
        "model": "",
        "prompt": "",
        "started_at": "",
        "completed_at": "",
        "duration_seconds": None,
        "result_preview": "",
    }


def text_sheet_row_search_haystack(item: dict[str, Any]) -> str:
    values = [
        item.get("record_id"),
        item.get("input_preview"),
        item.get("row"),
        item.get("job_id"),
        *((item.get("mapped_cells") or {}).values()),
    ]
    return " ".join(str(value or "") for value in values).lower()


def text_search_haystack(record: dict[str, Any]) -> str:
    metadata = record.get("metadata") or {}
    values = [
        metadata.get("record_id"),
        metadata.get("input_preview"),
        metadata.get("source_filename"),
        metadata.get("row_display_number"),
        record.get("job_id"),
    ]
    return " ".join(str(value or "") for value in values).lower()


def normalize_status_key(value: Any) -> str:
    raw = str(value or "queued").strip().lower().replace("-", "_")
    if raw in {"all", ""}:
        return "all"
    if raw in {"not_run", "not_started", "unprocessed"}:
        return "not_run"
    if raw in {"complete", "completed", "done"}:
        return "completed"
    if raw in {"failed", "error"}:
        return "failed"
    if raw in {"running", "rq_screening", "openai_running"}:
        return "running"
    return "queued"


def status_label(status_key: str) -> str:
    return {
        "not_run": "Not run",
        "completed": "Completed",
        "running": "Running",
        "queued": "Queued",
        "failed": "Failed",
    }.get(status_key, "Queued")


def text_row_order(metadata: dict[str, Any]) -> int:
    try:
        return int(metadata.get("project_row_order"))
    except (TypeError, ValueError):
        try:
            return int(metadata.get("row_display_number"))
        except (TypeError, ValueError):
            return 0


def next_project_row_order(project_id: str) -> int:
    max_order = -1
    for record in text_job_records(project_id):
        max_order = max(max_order, text_row_order(record.get("metadata") or {}))
    return max_order + 1


def export_columns_for_records(project_id: str, records: list[dict[str, Any]]) -> list[str]:
    columns: list[str] = []
    seen: set[str] = set()
    seen_sources: set[str] = set()
    ordered_source_ids: list[str] = []
    for record in records:
        source_id = str((record.get("metadata") or {}).get("source_id") or "")
        if source_id and source_id not in seen_sources:
            ordered_source_ids.append(source_id)
            seen_sources.add(source_id)
    for source_id in ordered_source_ids:
        if not source_id:
            continue
        manifest = read_json_if_exists(source_dir_for(project_id, source_id) / SOURCE_MANIFEST_FILENAME)
        for column in manifest.get("columns") or []:
            if column not in seen:
                columns.append(str(column))
                seen.add(str(column))
    if columns:
        return columns
    return ["record_id"]


def load_all_source_rows(project_id: str) -> dict[str, dict[int, dict[str, str]]]:
    source_rows: dict[str, dict[int, dict[str, str]]] = {}
    sources_root = projects.project_root(project_id) / SOURCE_FILES_DIRNAME
    if not sources_root.exists():
        return source_rows
    for source_dir in sources_root.iterdir():
        rows_path = source_dir / ROWS_JSONL_FILENAME
        if not rows_path.exists():
            continue
        rows_by_index: dict[int, dict[str, str]] = {}
        with rows_path.open("r", encoding="utf-8") as handle:
            for line in handle:
                if not line.strip():
                    continue
                payload = json.loads(line)
                rows_by_index[int(payload.get("source_row_index") or 0)] = payload.get("row") or {}
        source_rows[source_dir.name] = rows_by_index
    return source_rows


def read_text_if_exists(path: Path) -> str:
    if not path.exists():
        return ""
    return path.read_text(encoding="utf-8")


def read_json_if_exists(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    return jobs.read_json(path)


def queue_key(project_id: str, job_id: str) -> str:
    return f"{project_id}:{job_id}"


def job_id_from_key(key: str) -> str:
    return key.split(":", maxsplit=1)[1] if ":" in key else key


def stringify_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def is_blank_row(row: dict[str, str]) -> bool:
    return not any(str(value or "").strip() for value in row.values())


def safe_text_worksheet_name(name: str, used_names: set[str]) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]", "_", str(name or "Sheet")).strip() or "Sheet"
    cleaned = cleaned[:31]
    candidate = cleaned
    index = 1
    while candidate.lower() in used_names:
        suffix = f" ({index})"
        candidate = f"{cleaned[: max(1, 31 - len(suffix))]}{suffix}"
        index += 1
    used_names.add(candidate.lower())
    return candidate


def truncate(value: str, limit: int) -> str:
    text = str(value or "")
    if len(text) <= limit:
        return text
    return text[: max(0, limit - 1)].rstrip() + "..."


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()
